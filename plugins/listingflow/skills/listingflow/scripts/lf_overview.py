"""Decision-first operator views; never alters checks, facts, or approvals."""

from datetime import datetime

from lf_checklist import ITEMS, SHEETS, content_texts, report_current
from lf_experience import field_label, label
from lf_rules import content_blocks, digest


LAYOUT = "operator-overview-v1"
GOOD = {"PASS", "NOT_APPLICABLE"}
SUMMARY_HEADERS = ["审核结论", "当前结果", "说明", "负责方", "下一步", "查看详情"]
SUMMARY_WIDTHS = [23, 28, 47, 18, 38, 20]
ACTION_HEADERS = ["待办事项", "影响与原句", "处理建议", "负责方", "完成条件",
                  "处理意见", "确认人", "确认日期", "反馈状态"]
ACTION_WIDTHS = [24, 48, 38, 18, 38, 24, 12, 15, 16]


def configuration(item):
    return item["id"] == "bundle.approval" or (
        item["id"].startswith("R-") and item["status"] == "NEEDS_REVIEW" and item["path"] == "rules"
        and item.get("method") != "explicit_rule_review")


def local_time(value):
    if not value:
        return "未记录"
    return datetime.fromisoformat(value).astimezone().isoformat(sep=" ", timespec="seconds")


def decision(request):
    """All summaries and actionable lists share the same classification."""
    results = request["checks"]["results"]
    review = request.get("review") or {}
    texts = content_texts(request["content"])
    by_id = {}
    for result in results:
        by_id.setdefault(result["id"], []).append(result)
    current = report_current(request)
    required = {"semantic." + p for p, _ in content_blocks(request["content"])}
    required.update("checklist." + key for key, (_, _, severity) in ITEMS.items() if severity == "blocking")
    required.add("bundle.approval")
    if request.get("research_version"):
        required.add("research.current")
    missing = required - set(by_id)
    semantic_ok = current and all(
        by_id.get("semantic." + p) and all(r["status"] == "PASS" for r in by_id["semantic." + p])
        for p, _ in content_blocks(request["content"]))
    unresolved = [r for r in results if r["status"] not in GOOD]
    config = [r for r in unresolved if configuration(r)]
    actions, optional, handled, handled_fields = [], [], [], set()

    def add(title, status, finding, action, owner, done, target, kind, ids=()):
        item = {"title": title, "status": status, "finding": finding, "action": action,
                "owner": owner, "done": done, "target": target, "kind": kind, "ids": list(ids)}
        actions.append(item)
        return item

    if not current or missing:
        add("补齐本版审核", "审核未完成", "当前内容与审核未对应，或必要核查记录不完整。",
            "由助手针对当前内容重新检查，不沿用旧版结论。", "文案助手",
            "当前文案、规则和审核版本一致，必要检查齐全。", "内容检查", "review")

    fields = {f["field"] for f in request["facts"]["facts"]}
    disposition_ids = set()
    for row in request["facts"].get("reconciliation", []):
        if row["status"] == "adopted":
            continue
        disposition_ids.add("reliability.field." + row["field"])
        field_checks = by_id.get("reliability.field." + row["field"], [])
        field_failed = any(r["status"] == "FAIL" for r in field_checks)
        field_blocked = any(r["severity"] == "blocking" and r["status"] not in GOOD for r in field_checks)
        title = row.get("label", field_label(row["field"]))
        safely_omitted = (row["status"] in {"pending", "omitted"} and not row["critical"]
                          and not row["fact_ids"] and row["field"] not in fields and semantic_ok
                          and not field_failed and not field_blocked)
        if safely_omitted:
            handled_fields.add(row["field"])
            handled.append([title, "本版不采用，无需补填", row["reason"],
                            "仅在需要补写该字段时，由产品负责人确认批次、版本或口径。",
                            row["owner"], "来源依据"])
        else:
            finding = "\n".join(dict.fromkeys([row["reason"], *(r["reason"] for r in field_checks)]))
            add(title, "需修改" if field_failed else "待核实", finding, row["action"], row["owner"],
                "确认适用事实；或核实本版确未采用且不影响关键用途与条件。",
                "来源依据", "fix" if field_failed else "review", ["reliability.field." + row["field"]])

    catalog = {r["id"]: r for r in request.get("rule_catalog", [])}
    if request.get("research_version"):
        from lf_research import coverage, current_report
        if current_report(request):
            for key, item in coverage(request["research"]).items():
                if not item["complete"]:
                    attempts = request["research"]["channels"][key]
                    declined = all(a["status"] == "declined" for a in attempts)
                    if declined:
                        handled.append([item["name"], "按用户要求跳过", item["reason"],
                                        "用户要求补做时再采集。", "文案助手", "本地化自查表"])
                    else:
                        add(item["name"], item["status"], item["reason"],
                            "助手重试；如果页面要求店铺登录，由运营登录后继续。无需重新上传产品资料。",
                            "助手；登录由运营", "获得推荐词或明确记录本次可用范围。",
                            "本地化自查表", "research", ["research." + key])
    for item in unresolved:
        identifier = item["id"]
        if configuration(item) or identifier in disposition_ids:
            continue
        key = identifier.removeprefix("checklist.")
        entry = (review.get("operator_checklist") or {}).get(key, {})
        if request.get("research_version") and (
                identifier in {"checklist.local.search_data", "keywords.live_data"}
                or identifier.startswith("research.") and identifier != "research.current"):
            continue
        if (item.get("method") == "explicit_rule_review"
                and catalog.get(identifier, {}).get("kind") in {"image_source", "image_text"}):
            add("素材规则：" + catalog[identifier]["raw"], "素材待处理", item["reason"],
                "核对规则指向的素材与依据，不把专项规则失败写成文案风险词命中。",
                "规则/权利负责人", "相应素材符合已批准规则并重新完成专项核查。",
                "侵权词自查表", "asset", [identifier])
            continue
        if identifier in {"checklist.local.search_data", "keywords.live_data"} and item["status"] != "FAIL":
            if not any(r[0] == "实时搜索数据" for r in optional):
                optional.append(["实时搜索数据", "未验证，可选补充", item["reason"],
                                 "已有本站点关键词数据时再优化；无需为草稿临时补交。",
                                 "运营（可选）", "本地化自查表"])
            continue
        if key == "ip.visual":
            add("拟用图片与外观权利", "用图前确认", item["reason"],
                entry.get("action") or "确认拟刊登图片及相关权利材料。",
                "运营/权利负责人", "拟用素材、用途和必要授权得到相应负责人确认。",
                "侵权词自查表", "asset", [identifier])
            continue
        if item["severity"] == "info" and item["status"] != "FAIL":
            optional.append([ITEMS[key][1] if key in ITEMS else label(item["path"]),
                             "信息提示", item["reason"], "按具体范围决定是否补充；不冒充已完成核查。",
                             "文案助手", "规则与范围"])
            continue
        title = (ITEMS[key][1] if key in ITEMS else
                 "词库命中：" + catalog[identifier]["raw"] if identifier in catalog else label(item["path"]))
        paths = [e["path"] for e in entry.get("examples", [])] if key in ITEMS else [item["path"]]
        quotes = [f"{label(p)}：{texts[p]}" for p in paths if p in texts]
        finding = item["reason"]
        if quotes:
            # Keep all matched examples in the detail sheet; no arbitrary truncation.
            finding += "\n" + "\n".join(dict.fromkeys(quotes))
        target = ITEMS[key][0] if key in ITEMS else (
            "违禁词自查表" if catalog.get(identifier, {}).get("group") == "general_banned" else
            "侵权词自查表" if identifier in catalog else "内容检查")
        add(title, "需修改" if item["status"] == "FAIL" else "待复核",
            finding, entry.get("action") or "助手核对对应原句与依据，必要时修正并重查。",
            "文案助手", "对应问题修正或取得适用依据，本版重新检查。",
            target, "fix" if item["status"] == "FAIL" else "review", [identifier])
    if config:
        add("公司规则确认", "正式批准前确认",
            f"{len(config)}条配置或适用范围记录待确认，不等于文案有{len(config)}处违规。",
            "规则负责人统一确认公司规则与适用范围；运营不用逐条审批词库。",
            "规则负责人", "所需规则与范围明确批准，再按批准范围复查。",
            "规则与范围", "config", [r["id"] for r in config])

    approval = request.get("approval") or {}
    approval_current = (bool(approval) and approval.get("content_hash") == request["content_hash"]
                        and approval.get("checks_hash") == digest(request["checks"]))
    blockers = [r for r in unresolved if r["severity"] == "blocking"]
    asset_ids = {identifier for a in actions if a["kind"] == "asset" for identifier in a["ids"]}
    asset_blocked = any(r["id"] in asset_ids for r in blockers)
    if not approval_current:
        add("审阅并确认当前文案", "运营审阅", "当前文案尚无与本次检查对应的有效批准记录。",
            "核对标题、核心功能、供电和限制，回复保留或修改意见；无需替其他负责人签字。",
            "运营", "明确确认当前版本；正式批准仍需全部必要检查及规则条件满足。",
            "内容检查", "approval")
    kinds = {a["kind"] for a in actions}
    copy_status = ("需修改" if "fix" in kinds else
                   "审核未完成" if not current or missing else
                   "有事项待复核" if "review" in kinds else "本轮未发现必须修改项")
    formal = request["mode"] == "final" and approval_current and not blockers and current and not missing
    publish_status = ("已批准（仅本版文本范围）" if formal else
                      "仍为草稿，待规则确认" if config else
                      "仍为草稿，待素材规则确认" if asset_blocked else
                      "仍为草稿，先完成文案复核" if copy_status != "本轮未发现必须修改项" else
                      "仍为草稿，待运营确认" if not approval_current else "已确认文案，尚未正式导出")
    if request["mode"] == "final" and not formal:
        publish_status = "审批记录待核对"
    priority = {"fix": 0, "review": 1, "research": 2, "approval": 3, "config": 4, "asset": 5}
    actions.sort(key=lambda a: priority[a["kind"]])
    categories = []
    for sheet in SHEETS:
        keys = [k for k, (name, _, severity) in ITEMS.items() if name == sheet and severity == "blocking"]
        semantic = [r for k in keys for r in by_id.get("checklist." + k, [])]
        if sheet == SHEETS[2]:
            semantic.extend(r for r in results if r["id"] == "review.localization"
                            or r["id"].startswith("reliability.units."))
        elif sheet == SHEETS[3]:
            semantic.extend(r for r in results
                            if r["id"].startswith(("length.", "case.", "count.", "distinct.", "required.",
                                                    "semantic.", "editorial.", "reliability.", "keywords."))
                            and not r["id"].startswith("reliability.field.") and r["id"] != "keywords.live_data")
            semantic.extend(r for r in results if r["id"] in {"review.claims", "review.keywords"})
        elif sheet == SHEETS[1]:
            semantic.extend(by_id.get("checklist.compliance.ip", []))
        complete = all("checklist." + k in by_id for k in keys)
        groups = {"general_banned"} if sheet == SHEETS[0] else {"general_ip", "ebay_ip"} if sheet == SHEETS[1] else set()
        scans = [r for r in results if catalog.get(r["id"], {}).get("group") in groups]
        asset_scans = [r for r in scans if r.get("method") == "explicit_rule_review"
                       and catalog.get(r["id"], {}).get("kind") in {"image_source", "image_text"}]
        text_scans = [r for r in scans if r not in asset_scans]
        pending = sum(r["status"] not in GOOD for r in scans)
        failures = any(r["status"] == "FAIL" for r in semantic + text_scans)
        unknown = not current or not complete or any(r["status"] not in GOOD for r in semantic)
        result = "需修改" if failures else "尚待核对" if unknown else "本类文字已核对"
        if not failures and not unknown and groups and (pending or not scans):
            result = "文字已核对，词库待确认"
        if not failures and not unknown and any(r["status"] not in GOOD for r in asset_scans):
            result = "文字已核对，素材待处理"
        descriptions = {
            SHEETS[0]: "逐词扫描与上下文分别审核。",
            SHEETS[1]: "品牌、IP、赛事、奢侈品文字分别审核。",
            SHEETS[2]: "检查标题、五点、语气、单位和实际配置。",
            SHEETS[3]: "检查词频、夸大、跨类目、医疗暗示和侵权表达。",
        }
        description = descriptions[sheet]
        if groups:
            description += (f" {pending}条规则待确认。" if pending else
                            "当前无词库扫描记录。" if not scans else "本次适用规则已有结果。")
        external = "ip.visual" if sheet == SHEETS[1] else "local.search_data" if sheet == SHEETS[2] else None
        if external and (not by_id.get("checklist." + external)
                         or any(r["status"] not in GOOD for r in by_id.get("checklist." + external, []))):
            if external == "ip.visual":
                description += "图片与外观权利另核。"
            elif request.get("research_version"):
                description += "关键词来源见首页调研摘要及本地化明细。"
            else:
                description += "实时搜索数据未验证。"
        categories.append([sheet, result, description, "文案助手",
                           "异常见待办；需要时查看原句与依据。", sheet])
    return {"copy_status": copy_status, "publish_status": publish_status, "formal": formal,
            "actions": actions, "handled": handled, "optional": optional,
            "handled_fields": sorted(handled_fields), "categories": categories,
            "current": current, "missing_checks": sorted(missing)}


def tables(request):
    model = decision(request)
    actions = model["actions"]
    fix_or_review = any(a["kind"] in {"fix", "review"} for a in actions)
    next_step = ("先由助手处理文案问题，运营只回答资料无法解决的疑点。" if fix_or_review else
                 "审阅当前文案并回复意见；规则由负责人处理，用图前另核权利。")
    rows = [
        ["商品", "SKU " + request["sku"], f"{request['site']}站 / {request['language']} / 文案第{request['version']}版",
         "", "核对商品与站点是否正确。", "内容检查"],
        ["检查时间", local_time(request.get("checked_at", request["created"])), "本页只对应当前导出快照。",
         "", "修改文案后须重新核查。", "规则与范围"],
        ["文案检查", model["copy_status"], "包含事实、表达和四类自查；不以词库未命中数量替代质量判断。",
         "文案助手", "先处理文案异常。" if fix_or_review else "进入运营审阅，不等于已获发布批准。", "处理清单"],
        ["发布前确认", model["publish_status"], "文本批准与图片、商品权利核验分开；本工具不自动刊登。",
         "运营/相关负责人", "完成所属事项，不替其他负责人签字。", "处理清单"],
        ["你现在要做什么", "先处理下方待办" if actions else "本轮无新增待办",
         next_step if actions else "按已批准文本范围使用；新增资料或改文案后重查。",
         "运营", "已避开的非关键参数无需本轮补填。", "已处理与可选"],
        ["四类自查摘要", "文字核对与未核验范围分开", "", "", "", ""],
        *model["categories"],
        ["本轮待办", f"{len(actions)}项，下面展示优先事项", "全部事项、完成条件和反馈栏见处理清单。",
         "", "", "处理清单"],
    ]
    if request.get("research_version"):
        from lf_research import current_report, summary
        rows.insert(5, ["关键词调研", "已记录，范围见说明" if current_report(request) else "未完成或已过期",
                        summary(request["research"]) if current_report(request) else "由助手完成调研并绑定本版文案。",
                        "文案助手", "采集受限项见待办；词条及采用理由见明细。", "本地化自查表"])
    for item in actions[:3]:
        rows.append([item["title"], item["status"], item["done"], item["owner"], item["action"], "处理清单"])
    if len(actions) > 3:
        rows.append(["其余待办", f"另有{len(actions) - 3}项", "未从报告中删除，见完整处理清单。",
                     "", "不要将首页摘要当作全部已处理。", "处理清单"])
    if not actions:
        rows.append(["本轮待办", "无", "仅在本次检查及批准范围内成立。", "", "", "规则与范围"])
    rows.append(["已处理与可选", f"{len(model['handled'])}项本版不采用",
                 "产品差异仍保留；需要补写时再核实。可选数据不强制补交。",
                 "", "无需重新上传整套资料。", "已处理与可选"])
    rows.append(["反馈与批准", "填写意见不会自动批准",
                 "只确认自己负责的事项。回传意见后，由助手核对并导出新快照。",
                 "运营", "不要修改系统核查结果来消除待办。", "处理清单"])

    action_rows = [["填写说明", "填写后回传，系统结论不会自动改变。", "只填写蓝色反馈列。",
                    "", "助手复核并重新导出后，才更新结论。", "", "", "", ""]]
    for item in actions:
        action_rows.append([item["title"] + "\n" + item["status"], item["finding"], item["action"],
                            item["owner"], item["done"], "", "", "", ""])
    if not actions:
        action_rows.append(["本轮无待办", "仅限本次检查范围。", "新增内容后重新核查。",
                            "", "", "", "", "", ""])
    handled_rows = model["handled"] + model["optional"]
    if not handled_rows:
        handled_rows = [["本轮记录", "无", "没有另行搁置的字段或可选补充项。", "", "", "规则与范围"]]
    return [
        ("审核结论", SUMMARY_HEADERS, rows, SUMMARY_WIDTHS),
        ("处理清单", ACTION_HEADERS, action_rows, ACTION_WIDTHS),
        ("已处理与可选", ["事项", "本版处理", "保留的差异 / 范围", "何时需要再处理", "负责方", "查看依据"],
         handled_rows, [23, 29, 60, 48, 20, 20]),
    ]


def format_workbook(wb, name, request):
    """Navigation and accidental-edit protection; not identity or authorization."""
    if name != "自查" or request.get("report_layout") != LAYOUT:
        return
    from copy import copy
    from openpyxl.styles import Font, PatternFill, Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.hyperlink import Hyperlink

    red, amber, neutral = "FCE7E7", "FFF0C2", "EDF0F2"
    for ws in wb:
        if ws.title == "版本":
            continue
        ws.protection.sheet = True
        ws.protection.autoFilter = False
        ws.protection.selectLockedCells = False
        ws.protection.selectUnlockedCells = False
        ws.sheet_view.zoomScale = 90
        if ws.title == "审核结论":
            ws.auto_filter.ref = None
            ws.freeze_panes = "C2"
            ws.print_area = ws.dimensions
            ws.page_setup.fitToHeight = 1
            ws.sheet_properties.tabColor = "315447"
        for row in ws:
            for cell in row:
                value = str(cell.value or "")
                if value in wb.sheetnames and cell.row > 1:
                    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{value}'!A1",
                                               display=value, tooltip="跳转到" + value)
                    cell.font = Font(name="Microsoft YaHei", size=11, color="146B91", underline="single")
                if cell.column == 2 and cell.row > 1:
                    color = (red if value == "需修改" or value.startswith("✗") else
                             amber if any(s in value for s in ("待确认", "未完成", "待复核", "待核对", "待规则", "审批记录", "过期", "待处理")) else
                             neutral if any(s in value for s in ("可选", "不采用", "未检查", "未验证")) else None)
                    if color:
                        cell.fill = PatternFill("solid", fgColor=color)
        if ws.title == "审核结论":
            for row in ws:
                if row[0].value in {"四类自查摘要", "本轮待办"}:
                    for cell in row:
                        cell.fill = PatternFill("solid", fgColor="E4EDE8")
                        font = copy(cell.font)
                        font.bold = True
                        cell.font = font
        else:
            ws["A1"].hyperlink = Hyperlink(ref="A1", location="'审核结论'!A1", tooltip="返回审核结论")
        if ws.title == "处理清单":
            ws.freeze_panes = "B3"
            validation = DataValidation(type="list", formula1='"待处理,已回复,待复核,暂不采用"', allow_blank=True)
            validation.showErrorMessage = True
            validation.errorTitle = "请选择反馈状态"
            validation.error = "此处只记录反馈，不改变系统审核或批准状态。"
            ws.add_data_validation(validation)
            validation.add(f"I3:I{ws.max_row}")
            for row in ws.iter_rows(min_row=3, min_col=6, max_col=9):
                for cell in row:
                    cell.protection = Protection(locked=False)
                    cell.fill = PatternFill("solid", fgColor="EAF2FA")
    wb.active = 0
