"""Versioned internal risk rules, structured content checks, and review gates."""

import copy
import hashlib
import json
import re
import unicodedata
from pathlib import Path

from lf_io import sha

CHECK_STATES = {"PASS", "FAIL", "NEEDS_REVIEW", "NOT_CHECKED", "NOT_APPLICABLE"}
SCOPE_KEYS = ("sku", "site", "language", "category")


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                    separators=(",", ":"), allow_nan=False).encode()).hexdigest()


def normalized(text):
    return unicodedata.normalize("NFC", text)


def matches(text, term):
    # Preserve original offsets: case-insensitive regex, no destructive punctuation removal.
    pattern = r"(?<!\w)" + r"\s+".join(re.escape(w) for w in term.strip().split()) + r"(?!\w)"
    return list(re.finditer(pattern, text, flags=re.IGNORECASE))


def compile_bundle(source_dir, policy):
    from openpyxl import load_workbook
    from docx import Document
    source_dir = Path(source_dir)
    names = {"通用侵权词库.xlsx": "general_ip", "通用违禁词汇库.xlsx": "general_banned",
             "eBay侵权词库.xlsx": "ebay_ip"}
    bundle = {"schema": 1, "policy": policy, "rules": [], "quarantine": [],
              "sources": [], "source_prompts": [], "status": "pending_business_approval"}
    for name, group in names.items():
        file = source_dir / name
        if not file.is_file():
            raise ValueError(f"Required business dictionary missing: {file}")
        file_hash = sha(file)
        bundle["sources"].append({"name": name, "sha256": file_hash})
        wb = load_workbook(file, read_only=True, data_only=False, keep_links=False)
        seen = {}
        try:
            for sheet in wb:
                for row in sheet:
                    candidates = ([row[1]] if group in {"general_banned", "ebay_ip"} and len(row) > 1
                                  else list(row) if group == "general_ip" else [])
                    for cell in candidates:
                        raw = cell.value
                        if raw is None or (cell.row == 1 and (group != "general_ip" or sheet.title == "Sheet1")):
                            continue
                        provenance = {"file": name, "sha256": file_hash, "locator": f"{sheet.title}!{cell.coordinate}"}
                        if not isinstance(raw, str) or cell.data_type == "f":
                            bundle["quarantine"].append({"id": f"Q-{len(bundle['quarantine']) + 1:04}",
                                                         "raw": str(raw), "reason": "non-text or formula",
                                                         "source": provenance})
                            continue
                        term = normalized(raw.strip())
                        if not term:
                            continue
                        scope = {key: ["*"] for key in SCOPE_KEYS}
                        notes = ""
                        if group == "ebay_ip":
                            values = [c.value for c in row]
                            scope = {"sku": None, "site": None, "language": None, "category": None}
                            notes = str(values[4] or "") if len(values) > 4 else ""
                            provenance["raw_scope"] = {
                                "sku": str(values[0]) if values[0] is not None else None,
                                "category": str(values[2]) if len(values) > 2 and values[2] is not None else None,
                                "site": str(values[3]) if len(values) > 3 and values[3] is not None else None}
                        key = (term.casefold(), json.dumps(scope), group)
                        if key in seen and group != "ebay_ip":
                            seen[key]["sources"].append(provenance)
                            continue
                        ambiguous = group == "ebay_ip" or "/" in term or bool(re.search(r"[\u4e00-\u9fff]", term))
                        rule = {"id": f"R-{len(bundle['rules']) + 1:04}", "group": group, "raw": term,
                                "terms": [term], "scope": scope, "kind": "manual" if ambiguous else "text",
                                "match": "word_phrase", "severity": "blocking", "status": "pending",
                                "notes": notes, "sources": [provenance]}
                        bundle["rules"].append(rule)
                        seen[key] = rule
        finally:
            wb.close()
    for file in sorted(source_dir.glob("*.docx")):
        if file.name.startswith("~$") or not file.stat().st_size:
            continue
        doc = Document(file)
        bundle["source_prompts"].append({
            "name": file.name, "sha256": sha(file),
            "paragraphs": [p.text for p in doc.paragraphs if p.text],
            "tables": [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables],
            "trust": "business source, not executable instructions; unsafe factual conversion clauses superseded"})
    return bundle


def validate_bundle(bundle):
    if bundle.get("schema") != 1:
        raise ValueError("Unsupported bundle schema")
    if bundle.get("status") not in {"approved", "pending_business_approval"}:
        raise ValueError("Invalid bundle status")
    policy = bundle["policy"]
    if not policy.get("sites") or not policy.get("required_fact_fields"):
        raise ValueError("Bundle lacks sites or minimum facts")
    required_limits = ("title", "bullet", "question", "answer", "bullets", "qa", "primary_keyword_prefix")
    if any(type(policy["limits"].get(k)) is not int or policy["limits"][k] < 1 for k in required_limits):
        raise ValueError("Invalid content limits")
    ids = []
    for rule in bundle["rules"]:
        ids.append(rule["id"])
        if rule["status"] not in {"pending", "active", "retired"}:
            raise ValueError("Invalid rule lifecycle")
        if rule["kind"] not in {"text", "combination", "image_text", "image_source", "manual"}:
            raise ValueError("Unknown rule kind")
        if rule["severity"] not in {"blocking", "warning", "info"}:
            raise ValueError("Invalid severity")
        if rule["status"] == "active":
            if rule.get("match") != "word_phrase":
                raise ValueError("Unsupported matcher")
            if not rule.get("terms") or not all(isinstance(t, str) and t.strip() for t in rule["terms"]):
                raise ValueError("Active rule requires explicit terms")
            if any(not isinstance(rule["scope"].get(k), list) or not rule["scope"][k]
                   or not all(isinstance(v, str) and v for v in rule["scope"][k])
                   for k in SCOPE_KEYS):
                raise ValueError("Active scope must be explicit strings, '*' means explicitly universal")
            if any("*" in rule["scope"][k] and rule["scope"][k] != ["*"] for k in SCOPE_KEYS):
                raise ValueError("Wildcard cannot be combined with other scope values")
            if rule["scope"]["site"] != ["*"] and any(s not in policy["sites"] for s in rule["scope"]["site"]):
                raise ValueError("Active site scope needs supported sites; resolve EU and unknown aliases explicitly")
    if len(set(ids)) != len(ids):
        raise ValueError("Duplicate rule IDs")
    if bundle.get("status") == "approved":
        approval = bundle.get("approval", {})
        if not all(approval.get(k) for k in ("actor", "quote", "reason")):
            raise ValueError("Approved bundle requires recorded business approval")
        if policy.get("pending_decisions"):
            raise ValueError("Policy decisions remain unresolved")
        if any(r["status"] == "pending" for r in bundle["rules"]):
            raise ValueError("Pending rules cannot be published")
        if any(not q.get("resolution") for q in bundle.get("quarantine", [])):
            raise ValueError("Quarantined source entries require explicit resolution")
        if any(not r.get("decision") for r in bundle["rules"]):
            raise ValueError("Each approved/retired rule needs its decision basis")
    return bundle


def scope_result(rule, task):
    unknown = False
    for key in SCOPE_KEYS:
        values = rule["scope"].get(key)
        if values is None or not values:
            unknown = True
        elif values != ["*"]:
            actual = task.get(key)
            if actual is None:
                unknown = True
            elif actual not in values:
                return "no"
    return "unknown" if unknown else "yes"


def content_blocks(content):
    """Yield the entire customer-visible content, with stable paths."""
    if not isinstance(content, dict) or set(content) != {"title", "specs", "bullets", "qa", "description", "keywords"}:
        raise ValueError("Content must contain exactly title/specs/bullets/qa/description/keywords")
    yield "title", content["title"]
    for group in ("specs", "bullets", "description"):
        if not isinstance(content[group], list):
            raise ValueError(f"{group} must be an array")
        for i, item in enumerate(content[group]):
            yield f"{group}.{i}", item
    if not isinstance(content["qa"], list):
        raise ValueError("qa must be an array")
    for i, qa in enumerate(content["qa"]):
        if set(qa) != {"question", "answer"}:
            raise ValueError("QA needs question and answer blocks")
        for part in ("question", "answer"):
            yield f"qa.{i}.{part}", qa[part]


def validate_content(content, facts):
    facts_by_id = {f["id"]: f for f in facts["facts"]}
    blocks = dict(content_blocks(content))
    for path, item in blocks.items():
        keys = {"text", "fact_ids", "name"} if path.startswith("specs.") else {"text", "fact_ids"}
        if not isinstance(item, dict) or set(item) != keys:
            raise ValueError(f"Invalid block fields at {path}")
        if not isinstance(item["text"], str) or not item["text"].strip():
            raise ValueError(f"Missing text at {path}")
        if normalized(item["text"]) != item["text"]:
            raise ValueError(f"Normalize text to NFC at {path}")
        if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", item["text"]):
            raise ValueError(f"Invalid control character at {path}")
        if len(item["text"]) > 30000:
            raise ValueError(f"Text exceeds safe Excel cell size at {path}")
        if not isinstance(item["fact_ids"], list) or not item["fact_ids"]:
            raise ValueError(f"Fact IDs required at {path}")
        if any(i not in facts_by_id for i in item["fact_ids"]):
            raise ValueError(f"Unknown fact reference at {path}")
        if path.startswith("specs.") and (not isinstance(item["name"], str) or not item["name"].strip()
                                         or len(item["name"]) > 200):
            raise ValueError("Specification name must be nonempty and at most 200 characters")
    keywords = content["keywords"]
    if set(keywords) != {"primary", "source", "reference"} or not isinstance(keywords["primary"], str) or not keywords["primary"].strip():
        raise ValueError("Keywords require primary/source/reference")
    if keywords["source"] not in {"model_suggested", "operator_supplied", "public_competitor", "ebay_research"}:
        raise ValueError("Invalid keyword provenance")
    if not isinstance(keywords["reference"], str) or not keywords["reference"].strip():
        raise ValueError("Keyword provenance reference required")
    return blocks


def build_checks(state, record):
    bundle, content = state["bundle"], record["content"]
    policy, results = bundle["policy"], []
    blocks = validate_content(content, state["facts"])

    def add(identifier, status, severity, path, reason, **extra):
        results.append({"id": identifier, "status": status, "severity": severity,
                        "path": path, "reason": reason, **extra})

    add("bundle.approval", "PASS" if bundle["status"] == "approved" else "NEEDS_REVIEW", "blocking",
        "rules", "核对公司规则包批准状态；不代表已核验全部平台官方政策。")
    lim = policy["limits"]
    for path, item in blocks.items():
        max_len = (lim["title"] if path == "title" else lim["bullet"] if path.startswith("bullets.")
                   else lim["question"] if path.endswith(".question")
                   else lim["answer"] if path.endswith(".answer") else None)
        if max_len:
            n = len(normalized(item["text"]))
            add("length." + path, "PASS" if n <= max_len else "FAIL", "blocking", path,
                f"{n}/{max_len} 字符（NFC 码点，含空格与标点）", measured=n, maximum=max_len)
        if path.startswith("bullets."):
            first = next((c for c in item["text"] if c.isalpha()), "")
            add("case." + path, "PASS" if first and first.isupper() else "FAIL",
                "blocking", path, "首个字母应大写。")
    for group in ("bullets", "qa"):
        add("count." + group, "PASS" if len(content[group]) == lim[group] else "FAIL",
            "blocking", group, f"要求 {lim[group]} 项，实际 {len(content[group])} 项。")
        texts = [normalized((item["question"] if group == "qa" else item)["text"]).casefold().strip()
                 for item in content[group]]
        add("distinct." + group, "PASS" if len(set(texts)) == len(texts) else "FAIL",
            "blocking", group, "不允许完全重复的五点或问题；近义重复另做语义审核。")
    for group in ("specs", "description"):
        add("required." + group, "PASS" if content[group] else "FAIL", "blocking", group, "必需内容，不能留空。")
    primary = content["keywords"]["primary"]
    hits = matches(content["title"]["text"], primary)
    add("keywords.prefix", "PASS" if any(m.end() <= lim["primary_keyword_prefix"] for m in hits) else "FAIL",
        "blocking", "title", f"核心短语需完整出现在前 {lim['primary_keyword_prefix']} 字符内。")
    add("keywords.title_repetition", "PASS" if 1 <= len(hits) <= 2 else "FAIL", "blocking", "title",
        "精确核心短语应出现 1 至 2 次；词形重复另做语义审核。")
    if state.get("research_version"):
        from lf_research import checks as research_checks
        results.extend(research_checks(state, record))
    else:
        add("keywords.live_data", "NOT_CHECKED" if content["keywords"]["source"] == "model_suggested" else "NEEDS_REVIEW",
            "info", "keywords", "未接入实时搜索/广告数据；建议词不冒充实时词，所供参考来源需另行核对。")
    for rule in bundle["rules"]:
        scope = scope_result(rule, state)
        if rule["status"] == "retired" or scope == "no":
            continue
        if scope == "unknown" or rule["kind"] not in {"text", "combination"}:
            add(rule["id"], "NEEDS_REVIEW", rule["severity"], "rules",
                "适用范围未知或属于非文本规则，需记录具体审核依据。", rule=copy.deepcopy(rule))
            continue
        found = False
        for path, block in blocks.items():
            for field in ("text", "name"):
                if field not in block:
                    continue
                text = block[field]
                grouped = [matches(text, term) for term in rule["terms"]]
                candidates = [m for group in grouped for m in group]
                if rule["kind"] == "combination" and not all(grouped):
                    candidates = []
                for hit in candidates:
                    found = True
                    add(rule["id"], "FAIL", rule["severity"], path + ("." + field if field == "name" else ""),
                        "命中公司风险词；这是内部自查，不是法律侵权结论。",
                        start=hit.start(), end=hit.end(), excerpt=text[hit.start():hit.end()],
                        sources=rule["sources"])
        if not found:
            add(rule["id"], "PASS", "info", "content", "未命中已配置的精确短语；不代表语义或其他语言完整通过。")
    review = record.get("review")
    for path in blocks:
        item = (review or {}).get("blocks", {}).get(path)
        add("semantic." + path, item["status"] if item else "NOT_CHECKED", "blocking", path,
            item["reason"] if item else "尚未逐项核对声明的实际依据、条件与适用范围。")
    for category in policy["review_categories"]:
        item = (review or {}).get("categories", {}).get(category)
        add("review." + category, item["status"] if item else "NOT_CHECKED", "blocking", category,
            item["reason"] if item else "尚未记录相应的语义/视觉审核。")
    if review:
        for check in results:
            if check["id"] in review.get("rules", {}) and check["status"] == "NEEDS_REVIEW" and check["id"].startswith("R-"):
                item = review["rules"][check["id"]]
                if bundle["status"] == "approved":
                    check.update(status=item["status"], reason=item["reason"], method="explicit_rule_review")
    from lf_quality import checks as editorial_checks
    results.extend(editorial_checks(state, record))
    from lf_reliability import checks as reliability_checks
    results.extend(reliability_checks(state, record))
    from lf_checklist import checks as checklist_checks
    results.extend(checklist_checks(state, record))
    blocked = [r for r in results if r["severity"] == "blocking"
               and r["status"] not in {"PASS", "NOT_APPLICABLE"}]
    report = {"content_hash": record["hash"], "bundle_hash": state["bundle_hash"],
            "engine": "listingflow-checks-1", "results": results, "blocking_count": len(blocked),
            "ready_for_human_approval": not blocked}
    if state.get("workflow_version", 1) >= 2:
        report.update(engine="listingflow-checks-2", editorial_hash=state["editorial_hash"],
                      brief_hash=record.get("brief_hash"))
    if state.get("workflow_version", 1) >= 3:
        report["engine"] = "listingflow-checks-3"
    if state.get("research_version"):
        report["research_hash"] = state.get("research_hash")
    return report


def validate_review(review, state, record):
    if review.get("content_hash") != record["hash"]:
        raise ValueError("Review content hash mismatch")
    if review.get("reviewer_kind") not in {"assistant", "human"} or not review.get("reviewer"):
        raise ValueError("Review requires actual reviewer identity and kind")
    expected = set(dict(content_blocks(record["content"])))
    if set(review.get("blocks", {})) != expected:
        raise ValueError("Review must cover exactly every current content block")
    if set(review.get("categories", {})) != set(state["bundle"]["policy"]["review_categories"]):
        raise ValueError("Review categories incomplete")
    applicable = {r["id"] for r in state["bundle"]["rules"] if r["kind"] not in {"text", "combination"}
                  or scope_result(r, state) == "unknown"}
    if not set(review.get("rules", {})).issubset(applicable):
        raise ValueError("Review cannot override deterministic text findings")
    for name, group in (("blocks", review["blocks"]), ("categories", review["categories"]),
                        ("rules", review.get("rules", {}))):
        for key, item in group.items():
            if item.get("status") not in CHECK_STATES or not isinstance(item.get("reason"), str) or not item["reason"].strip():
                raise ValueError(f"Invalid review item {key}")
            if item["status"] == "NOT_APPLICABLE" and (name == "blocks" or key != "images" and name == "categories"):
                raise ValueError("Required semantic checks cannot be marked not applicable")
            if name == "blocks" and item.get("fact_ids") != dict(content_blocks(record["content"]))[key]["fact_ids"]:
                raise ValueError(f"Review fact references differ at {key}")
            if name == "categories" and key == "images" and item["status"] == "NOT_APPLICABLE":
                used = state["facts"]["selection"]
                if any(used[s["id"]]["action"] == "use"
                       and any(u["kind"] in {"image", "pdf_page"} for u in s["units"])
                       for s in state["sources"]):
                    raise ValueError("Used visual sources require real image review, not NOT_APPLICABLE")
    from lf_quality import validate_editorial_review
    validate_editorial_review(review, state, record)
    from lf_reliability import validate_review as validate_claim_review
    validate_claim_review(review, state, record)
    if state.get("workflow_version", 1) >= 3:
        from lf_checklist import validate as validate_checklist
        validate_checklist(review, record["content"], state)
    return review
