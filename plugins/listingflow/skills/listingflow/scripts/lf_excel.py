"""Deterministic Excel serialization after native artifact export failed on Windows."""

import math
from pathlib import Path


def tables(request):
    content = request["content"]
    state = "已审核版本" if request["mode"] == "final" else "草稿：未通过正式导出审核"
    metadata = [
        ["mode", state], ["batch", request["batch"]], ["SKU", request["sku"]], ["site", request["site"]],
        ["language", request["language"]], ["version", request["version"]], ["created_utc", request["created"]],
        ["content_hash", request["content_hash"]], ["bundle_hash", request["bundle_hash"]], ["facts_hash", request["facts_hash"]],
        ["rule_scope", "公司规则自查；不代表完整平台政策核验或法律结论"],
    ]
    if request.get("workflow_version", 1) >= 3:
        from lf_operator import tables as operator_tables
        metadata.extend([["editorial_hash", request["editorial_hash"]], ["brief_hash", request["brief_hash"]]])
        return operator_tables(request, metadata)
    rows = [["状态", state, ""], ["SKU", request["sku"], ""], ["国家/站点", request["site"], ""],
            ["语言", request["language"], ""], ["标题", content["title"]["text"], ", ".join(content["title"]["fact_ids"])]]
    for i, block in enumerate(content["specs"], 1):
        rows.append([f"规格 {i}：{block['name']}", block["text"], ", ".join(block["fact_ids"])])
    for i, block in enumerate(content["bullets"], 1):
        rows.append([f"五点 {i}", block["text"], ", ".join(block["fact_ids"])])
    for i, qa in enumerate(content["qa"], 1):
        for key, label in (("question", "问题"), ("answer", "回答")):
            rows.append([f"{label} {i}", qa[key]["text"], ", ".join(qa[key]["fact_ids"])])
    for i, block in enumerate(content["description"], 1):
        rows.append([f"详情 {i}", block["text"], ", ".join(block["fact_ids"])])
    rows.extend([["内容版本", request["version"], ""], ["生成时间 UTC", request["generated"], ""],
                 ["导出时间 UTC", request["created"], ""]])
    severity = {"blocking": "阻断", "warning": "警告", "info": "提示"}
    names = {"PASS": "通过", "FAIL": "失败", "NEEDS_REVIEW": "待确认", "NOT_CHECKED": "未检查", "NOT_APPLICABLE": "不适用"}
    check_rows = [[request["batch"], state, "", "", f"阻断项：{request['checks']['blocking_count']}", "", ""]]
    for item in request["checks"]["results"]:
        sources = item.get("sources", item.get("rule", {}).get("sources", []))
        check_rows.append([item["id"], names[item["status"]], severity[item["severity"]], item["path"],
                           item["reason"], item.get("excerpt", ""),
                           "\n".join(f"{s['file']} {s['locator']}" for s in sources)])
    evidence = []
    for fact in request["facts"]["facts"]:
        evidence.append(["事实", fact["id"], fact["field"], fact["value"], "\n".join(fact["sources"]), fact["support_reason"]])
    referenced = {sid for fact in request["facts"]["facts"] for sid in fact["sources"]}
    for source in request["sources"]:
        selection = request["facts"]["selection"][source["id"]]
        evidence.append(["资料", source["id"], source["name"], source["status"], source.get("sha256", ""), selection["reason"]])
        for unit in source.get("units", []):
            if unit["id"] in referenced:
                evidence.append(["来源片段", unit["id"], source["name"], unit["locator"], unit["text"], "原始提取；视觉观察另列"])
        for gap in source["gaps"]:
            d = request["facts"]["coverage"].get(gap["id"])
            evidence.append(["覆盖缺口", gap["id"], source["name"], gap["locator"], gap["reason"],
                             f"{d['action']}: {d['reason']}" if d else f"{selection['action']}: {selection['reason']}"])
    for observation in request.get("observations", []):
        evidence.append(["观察", observation["id"], observation["unit_id"], observation["method"],
                         observation["text"], f"{observation['actor']}: {observation['reason']}"])
    result = {
        "文案": [("文案", ["字段", "内容", "事实编号"], rows, [30, 88, 26]),
                 ("版本", ["字段", "值"], metadata, [24, 90])],
        "自查": [("自查", ["检查项", "状态", "级别", "内容位置", "依据/处理说明", "命中片段", "规则来源"],
                 check_rows, [33, 24, 15, 28, 72, 27, 45]),
                 ("依据", ["类型", "编号", "字段/文件", "值/位置", "来源/缺口", "核对/处置"], evidence, [16, 22, 32, 52, 65, 70]),
                 ("版本", ["字段", "值"], metadata, [24, 90])],
    }
    if request.get("workflow_version", 1) >= 2:
        metadata.extend([["editorial_hash", request["editorial_hash"]], ["brief_hash", request["brief_hash"]]])
        review = request.get("review") or {}
        for fid, excerpts in review.get("fact_evidence", {}).items():
            for excerpt in excerpts:
                evidence.append(["审核摘录", fid, excerpt["source_id"], excerpt.get("observation_id", ""),
                                 excerpt["quote"], "摘录存在不等于语义支持；须结合来源和本版审核判断。"])
        quality_labels = {"differentiation": "卖点互补", "buyer_usefulness": "购买决策价值",
                          "naturalness": "语言自然", "conciseness": "表达简洁",
                          "keyword_fit": "关键词相关", "conditions": "条件完整"}
        for dimension, item in review.get("quality", {}).items():
            examples = "\n".join(f"{e['path']}: {e['quote']}" for e in item.get("examples", []))
            evidence.append(["文案审核", quality_labels[dimension], names[item["status"]],
                             item["reason"], examples, f"{review.get('reviewer', '')} / {review.get('reviewer_kind', '')}"])
        operator = request["operator"]
        action_rows = [["当前阶段", operator["next_owner"], "", operator["stage"]],
                       ["下一步", operator["next_owner"], "", operator["next_action"]]]
        for group in operator["issues"]:
            detail = group["action"]
            detail += f"\n本组 {group['count']} 项，其中阻断 {group['blocking']} 项、警告 {group['warnings']} 项。"
            detail += "\n" + "\n".join(f"{e['path']}：{e['reason']}" for e in group["examples"])
            action_rows.append([{"business_rules": "规则待确认", "fix_copy": "文案需修改",
                                 "review_copy": "审核与质量提示", "keyword_data": "关键词来源说明",
                                 "other": "其他问题"}[group["group"]], group["owner"], group["count"], detail])
        if not operator["issues"]:
            action_rows.append(["当前检查", "", 0, "没有未处理检查项；不代表法律或平台政策的全面保证。"])
        result["自查"].insert(0, ("处理清单", ["事项", "负责方", "检查项数", "说明与下一步"], action_rows, [25, 19, 14, 110]))
    return result


def write_pair(request, output):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    output = Path(output)
    for name, sheets in tables(request).items():
        wb = Workbook()
        wb.remove(wb.active)
        for title, headers, rows, widths in sheets:
            ws = wb.create_sheet(title)
            ws.sheet_view.showGridLines = False
            for col, width in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(col)].width = width
            for r, row in enumerate([headers, *rows], 1):
                lines = 1
                for col, value in enumerate(row, 1):
                    if isinstance(value, str) and len(value.encode("utf-16-le")) // 2 > 32700:
                        raise ValueError("Excel cell exceeds safe capacity; split text instead of truncating")
                    cell = ws.cell(r, col, value)
                    if isinstance(value, str):
                        cell.data_type = "s"
                        cell.number_format = "@"
                    cell.alignment = Alignment(vertical="top", wrap_text=True,
                                               horizontal="left" if isinstance(value, str) else "right")
                    cell.font = Font(name="Microsoft YaHei", size=11, color="202420")
                    if r == 1:
                        cell.fill = PatternFill("solid", fgColor="315447")
                        cell.font = Font(name="Microsoft YaHei", size=11, color="FFFFFF", bold=True)
                    visual_lines = sum(max(1, math.ceil(sum(2 if ord(c) > 255 else 1 for c in part) /
                                                        max(8, widths[col - 1] - 3)))
                                       for part in str(value or "").split("\n"))
                    lines = max(lines, visual_lines)
                ws.row_dimensions[r].height = 30 if r == 1 else min(409, max(27, lines * 17 + 9))
            if len(rows) > 11:
                ws.freeze_panes = "A2"
            if request.get("workflow_version", 1) >= 3:
                ws.freeze_panes = "B2"
                ws.auto_filter.ref = ws.dimensions
                ws.sheet_view.zoomScale = 85
                if title == "版本":
                    ws.sheet_state = "hidden"
                if title.endswith("自查表"):
                    ws.freeze_panes = "C2"
                    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
                        cell = row[0]
                        value = str(cell.value or "")
                        color = ("FCE7E7" if value.startswith("✗") else
                                 "FFF0C2" if value in {"待确认", "未检查"} else None)
                        if color:
                            cell.fill = PatternFill("solid", fgColor=color)
                    ws.sheet_properties.tabColor = "315447"
                if title == "处理清单" and request.get("report_layout") != "operator-overview-v1":
                    validation = DataValidation(type="list", formula1='"待处理,已回复,待复核,本版不采用"', allow_blank=True)
                    validation.errorTitle = "请选择处理结果"
                    validation.error = "此处仅记录反馈，不会自动批准文案。"
                    validation.showErrorMessage = True
                    ws.add_data_validation(validation)
                    validation.add(f"H5:H{max(5, ws.max_row)}")
                    for row in ws.iter_rows(min_row=5, min_col=6, max_col=8):
                        for cell in row:
                            cell.fill = PatternFill("solid", fgColor="EAF2FA")
                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        if cell.value in {"需修改", "待核实", "未检查", "待核实，本版不写", "配置待确认，不等于文案违规"}:
                            cell.fill = PatternFill("solid", fgColor="FFF0C2")
            if title == "自查":
                for text, color in (("失败", "FCE7E7"), ("待确认", "FFF0C2"), ("未检查", "FFF0C2")):
                    ws.conditional_formatting.add(f"B2:B{len(rows) + 1}",
                        CellIsRule(operator="equal", formula=[f'"{text}"'], fill=PatternFill("solid", fgColor=color)))
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.orientation = "landscape"
            ws.page_setup.paperSize = ws.PAPERSIZE_A3
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.print_title_rows = "1:1"
        if request.get("report_layout") == "operator-overview-v1":
            from lf_overview import format_workbook
            format_workbook(wb, name, request)
        wb.save(output / f"{name}.xlsx")
        wb.close()
    verify_pair(request, output)


def verify_pair(request, output):
    from openpyxl import load_workbook
    for name, sheets in tables(request).items():
        wb = load_workbook(Path(output) / f"{name}.xlsx", read_only=True, data_only=False)
        try:
            if wb.sheetnames != [s[0] for s in sheets]:
                raise ValueError("Export sheet mismatch")
            for title, headers, rows, _ in sheets:
                actual = list(wb[title].iter_rows())
                expected = [headers, *rows]
                if len(actual) != len(expected):
                    raise ValueError("Export row count mismatch")
                for row, want in zip(actual, expected):
                    if len(row) != len(want):
                        raise ValueError("Export column count mismatch")
                    for cell, value in zip(row, want):
                        if cell.data_type == "f":
                            raise ValueError("Unexpected executable formula in text export")
                        if (cell.value if cell.value is not None else "") != (value if value is not None else ""):
                            raise ValueError(f"Export value mismatch: {title}!{cell.coordinate}")
        finally:
            wb.close()
