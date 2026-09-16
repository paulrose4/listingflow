"""Behavioral tests using synthetic inputs. No real product approvals are created."""

import copy
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import listingflow as lf
from lf_io import xml
from lf_rules import (build_checks, content_blocks, digest, matches, scope_result,
                      validate_bundle)


def write_xlsx_fixture(path, cells, dimension=None):
    """Minimal OOXML test fixture, not a user-authored workbook."""
    rows = {}
    for coord, value, typ, fmt in cells:
        row = "".join(c for c in coord if c.isdigit())
        if typ == "formula":
            body = f'<f>{escape(value)}</f>'
            attrs = ""
        elif typ == "number":
            body, attrs = f"<v>{value}</v>", ""
        else:
            body, attrs = f"<is><t>{escape(value)}</t></is>", ' t="inlineStr"'
        rows.setdefault(row, []).append(f'<c r="{coord}"{attrs} s="{fmt}">{body}</c>')
    data = "".join(f'<row r="{r}">{"".join(values)}</row>' for r, values in rows.items())
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/>'
                   '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
                   '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                   '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/></Types>')
        z.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        z.writestr("xl/workbook.xml", '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                   '<sheets><sheet name="Product" sheetId="1" r:id="rId1"/></sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
                   '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>')
        z.writestr("xl/styles.xml", '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                   '<numFmts count="1"><numFmt numFmtId="164" formatCode="000000"/></numFmts>'
                   '<fonts count="1"><font><sz val="11"/><name val="Arial"/></font></fonts>'
                   '<fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills>'
                   '<borders count="1"><border/></borders><cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
                   '<cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
                   '<xf numFmtId="164" fontId="0" fillId="0" borderId="0" xfId="0"/></cellXfs>'
                   '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles></styleSheet>')
        z.writestr("xl/worksheets/sheet1.xml", '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                   + (f'<dimension ref="{dimension}"/>' if dimension else "")
                   + f"<sheetData>{data}</sheetData></worksheet>")


def fixture_content():
    def block(text, *facts):
        return {"text": text, "fact_ids": list(facts or ("F1",))}
    return {
        "title": block("Desk Organizer with Three Compartments", "F1", "F3"),
        "specs": [{"name": "Compartments", **block("3", "F3")},
                  {"name": "Dimensions", **block("18 x 10 x 9 cm", "F4")}],
        "bullets": [
            block("Three compartments separate pens and small desk items", "F2", "F3"),
            block("An 18 x 10 cm footprint fits the stated desktop space", "F4"),
            block("Open compartments keep stored items accessible", "F5"),
            block("Polypropylene construction for everyday desk organization", "F6"),
            block("Wipe clean with a damp cloth as needed", "F7"),
        ],
        "qa": [
            {"question": block("Can it hold pens?", "F2"), "answer": block("Yes, it is designed to organize pens and small desk items.", "F2")},
            {"question": block("Does it have three compartments?", "F3"), "answer": block("Yes, it has three separate compartments.", "F3")},
            {"question": block("How large is this desk organizer?", "F4"), "answer": block("It measures 18 x 10 x 9 cm.", "F4")},
            {"question": block("Are the compartments open?", "F5"), "answer": block("Yes, the compartments have open tops.", "F5")},
            {"question": block("How should I clean it?", "F7"), "answer": block("Wipe it with a damp cloth.", "F7")},
        ],
        "description": [block("This desk organizer separates pens and small desk items across three open compartments.", "F1", "F2", "F3", "F5")],
        "keywords": {"primary": "Desk Organizer", "source": "model_suggested",
                     "reference": "Synthetic test suggestion; no live keyword data."},
    }


def approved_bundle():
    policy = lf.load(lf.SKILL / "assets/policy.json")
    policy.update(status="approved", pending_decisions=[])
    return {"schema": 1, "status": "approved", "policy": policy, "rules": [], "quarantine": [],
            "approval": {"actor": "SYNTHETIC TEST ONLY", "quote": "Fixture approval only", "reason": "Not a business approval"},
            "sources": [], "source_prompts": []}


def fixture_review(state):
    record = state["contents"][-1]
    return {"content_hash": record["hash"], "reviewer": "SYNTHETIC TEST REVIEW",
            "reviewer_kind": "human",
            "blocks": {p: {"status": "PASS", "fact_ids": b["fact_ids"], "reason": "Synthetic fixture claim checked"}
                       for p, b in content_blocks(record["content"])},
            "categories": {k: {"status": "NOT_APPLICABLE" if k == "images" else "PASS",
                               "reason": "Synthetic text-only fixture verification"}
                           for k in state["bundle"]["policy"]["review_categories"]},
            "rules": {}}


class FlowCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="listingflow-test-")
        self.root = Path(self.tmp.name)
        self.source = self.root / "sources"
        self.source.mkdir()
        self.bundle = self.root / "test-bundle.json"
        lf.write_new(self.bundle, approved_bundle())
        result = lf.run(lf.parser().parse_args([
            "init", "--root", str(self.root / "tasks"), "--source-root", str(self.source),
            "--sku", "000123", "--site", "US", "--language", "en", "--bundle", str(self.bundle)]))
        self.task = Path(result["task"])
        self.index = 0
        if getattr(self, "use_v3", False):
            pass
        elif getattr(self, "use_v2", False):
            state = lf.current(self.task)
            state.pop("research_version", None)
            state["workflow_version"] = 2
            lf.commit(self.task, state["revision"], "v2-test-fixture", state)
        else:
            # Preserve these original regression cases as saved v1 task fixtures.
            state = lf.current(self.task)
            for key in ("workflow_version", "editorial", "editorial_hash", "research_version"):
                state.pop(key, None)
            lf.commit(self.task, state["revision"], "legacy-test-fixture", state)

    def tearDown(self):
        self.tmp.cleanup()

    def command(self, name, *extra):
        args = [name, "--task", str(self.task)]
        if name in {"ingest", "observe", "statement", "freeze", "save-content", "review", "check", "approve", "export",
                    "brief", "revise", "prepare", "lock", "research"}:
            args += ["--expect", str(lf.current(self.task)["revision"])]
        return lf.run(lf.parser().parse_args([*args, *extra]))

    def payload(self, value):
        self.index += 1
        path = self.root / f"payload-{self.index}.json"
        lf.write_new(path, value)
        return str(path)

    def ingested(self):
        p = self.source / "中文产品资料.xlsx"
        values = ["000123", "Desk organizer", "Organizes pens and small desk items", "3",
                  "18 x 10 x 9 cm", "Open tops", "Polypropylene", "Wipe with damp cloth"]
        write_xlsx_fixture(p, [(f"A{i + 1}", text, "text", 0) for i, text in enumerate(values)])
        self.command("ingest", "--path", str(p))
        return lf.current(self.task)

    def facts_payload(self):
        state = lf.current(self.task)
        source = state["sources"][0]
        by_text = {u["text"]: u["id"] for u in source["units"]}
        vals = [("product_name", "Desk organizer"), ("core_function", "Organizes pens and small desk items"),
                ("compartments", "3"), ("dimensions", "18 x 10 x 9 cm"), ("open_tops", "Open tops"),
                ("material", "Polypropylene"), ("care", "Wipe with damp cloth")]
        return {"sku": "000123", "site": "US", "product_revision": "Synthetic fixture revision 1",
                "selection": {source["id"]: {"action": "use", "sku": "000123", "site": "US",
                                             "reason": "Synthetic fixture explicitly labels target SKU"}},
                "coverage": {}, "required_fields": [], "unresolved": [], "decisions": [],
                "facts": [{"id": f"F{i + 1}", "field": field, "value": value, "sku": "000123",
                           "status": "confirmed", "sources": [by_text[value]], "support_reason": "Explicit synthetic source value"}
                          for i, (field, value) in enumerate(vals)]}

    def ready_content(self):
        self.ingested()
        self.command("freeze", "--json", self.payload(self.facts_payload()))
        return self.command("save-content", "--json", self.payload(fixture_content()))

    def ready_review(self):
        self.ready_content()
        self.command("review", "--json", self.payload(fixture_review(lf.current(self.task))))
        return lf.current(self.task)

    def test_no_english_draft_required(self):
        self.ready_content()
        self.assertEqual(len(lf.current(self.task)["contents"]), 1)

    def test_all_rows_beyond_1000_and_inflated_dimension(self):
        p = self.source / "late.xlsx"
        write_xlsx_fixture(p, [("A1", "SKU", "text", 0), ("B1501", "000123", "text", 0)], "A1:XFD1048576")
        self.command("ingest", "--path", str(p))
        self.assertTrue(any(u["text"] == "000123" for u in lf.current(self.task)["sources"][0]["units"]))

    def test_numeric_leading_zero_display_preserved(self):
        p = self.source / "numeric.xlsx"
        write_xlsx_fixture(p, [("A1", 123, "number", 1)])
        self.command("ingest", "--path", str(p))
        u = lf.current(self.task)["sources"][0]["units"][1]
        self.assertEqual((u["text"], u["raw_value"]), ("000123", 123))

    def test_formula_missing_cache_not_fact(self):
        p = self.source / "formula.xlsx"
        write_xlsx_fixture(p, [("A1", "1+1", "formula", 0)])
        self.command("ingest", "--path", str(p))
        s = lf.current(self.task)["sources"][0]
        self.assertEqual(s["status"], "PARTIAL")
        self.assertEqual(s["units"][1]["text"], "")
        self.assertTrue(s["gaps"])

    def test_empty_and_unsupported_inventory(self):
        (self.source / "~$locked.xlsx").touch()
        (self.source / "video.mp4").write_bytes(b"test")
        self.command("ingest", "--path", str(self.source))
        self.assertEqual({s["status"] for s in lf.current(self.task)["sources"]}, {"excluded", "unsupported"})

    def test_out_of_scope_rejected(self):
        p = self.root / "outside.xlsx"
        p.write_bytes(b"test")
        with self.assertRaisesRegex(ValueError, "outside authorized"):
            self.command("ingest", "--path", str(p))

    def test_revision_conflict(self):
        self.ingested()
        stale = lf.current(self.task)
        lf.commit(self.task, stale["revision"], "other-session", stale)
        with self.assertRaisesRegex(ValueError, "Revision conflict"):
            lf.commit(self.task, stale["revision"], "stale", stale)

    def test_missing_fact_source_rejected(self):
        self.ingested()
        p = self.facts_payload()
        p["facts"][0]["sources"] = ["S9999:U1"]
        with self.assertRaisesRegex(ValueError, "source IDs"):
            self.command("freeze", "--json", self.payload(p))

    def test_cross_sku_rejected(self):
        self.ingested()
        p = self.facts_payload()
        p["facts"][0]["sku"] = "other"
        with self.assertRaisesRegex(ValueError, "applicable"):
            self.command("freeze", "--json", self.payload(p))

    def test_unresolved_conflict_rejected(self):
        self.ingested()
        p = self.facts_payload()
        p["unresolved"] = ["voltage conflict"]
        with self.assertRaisesRegex(ValueError, "Unresolved"):
            self.command("freeze", "--json", self.payload(p))

    def test_duplicate_field_conflict_rejected(self):
        self.ingested()
        p = self.facts_payload()
        f = copy.deepcopy(p["facts"][0])
        f.update(id="F8", value="different product")
        p["facts"].append(f)
        with self.assertRaisesRegex(ValueError, "Conflicting"):
            self.command("freeze", "--json", self.payload(p))

    def test_excluded_source_cannot_support_fact(self):
        self.ingested()
        p = self.facts_payload()
        p["selection"]["S0001"]["action"] = "exclude"
        with self.assertRaisesRegex(ValueError, "excluded source"):
            self.command("freeze", "--json", self.payload(p))

    def test_frozen_facts_no_new_ingestion(self):
        self.ready_content()
        with self.assertRaisesRegex(ValueError, "frozen"):
            self.command("ingest", "--path", str(self.source))

    def test_partial_edit_protects_other_blocks(self):
        self.ready_content()
        p = fixture_content()
        p["title"]["text"] = "Changed Desk Organizer"
        with self.assertRaisesRegex(ValueError, "Unselected block"):
            self.command("save-content", "--json", self.payload(p), "--only", "bullets.2")

    def test_partial_edit_creates_version_and_invalidates_review(self):
        self.ready_review()
        p = fixture_content()
        p["bullets"][2]["text"] = "Open tops make the stored items accessible"
        self.command("save-content", "--json", self.payload(p), "--only", "bullets.2")
        s = lf.current(self.task)
        self.assertIsNone(s["contents"][-1]["review"])
        self.assertIsNotNone(s["contents"][0]["review"])
        diff = self.command("compare", "--first", "1", "--second", "2")
        self.assertEqual([c["path"] for c in diff["changes"]], ["bullets.2"])

    def test_missing_semantic_review_blocks_approval(self):
        self.ready_content()
        with self.assertRaisesRegex(ValueError, "Blocking"):
            self.command("approve", "--json", self.payload({"content_hash": lf.current(self.task)["contents"][-1]["hash"],
                                                            "actor": "fixture", "quote": "fixture only"}))

    def test_review_bound_to_content_hash(self):
        self.ready_content()
        review = fixture_review(lf.current(self.task))
        review["content_hash"] = "stale"
        with self.assertRaisesRegex(ValueError, "hash mismatch"):
            self.command("review", "--json", self.payload(review))

    def test_incomplete_review_rejected(self):
        self.ready_content()
        review = fixture_review(lf.current(self.task))
        del review["blocks"]["title"]
        with self.assertRaisesRegex(ValueError, "exactly every"):
            self.command("review", "--json", self.payload(review))

    def test_complete_review_then_approval(self):
        self.ready_review()
        content_hash = lf.current(self.task)["contents"][-1]["hash"]
        self.command("approve", "--json", self.payload({"content_hash": content_hash, "actor": "TEST", "quote": "Synthetic approval only"}))
        self.assertTrue(lf.current(self.task)["contents"][-1]["approval"])

    def test_snapshot_tamper_blocks_resume(self):
        self.ingested()
        s = lf.current(self.task)
        (self.task / s["sources"][0]["snapshot"]).write_bytes(b"tampered")
        with self.assertRaisesRegex(ValueError, "snapshot missing or modified"):
            self.command("status")

    def test_invalid_content_reference_rejected(self):
        self.ready_content()
        p = fixture_content()
        p["title"]["fact_ids"] = ["F999"]
        with self.assertRaisesRegex(ValueError, "Unknown fact"):
            self.command("save-content", "--json", self.payload(p))

    def test_long_title_blocks(self):
        self.ready_content()
        p = fixture_content()
        p["title"]["text"] = "Desk Organizer " + "x" * 80
        self.command("save-content", "--json", self.payload(p))
        checks = self.command("show", "--section", "checks")
        self.assertTrue(any(r["id"] == "length.title" and r["status"] == "FAIL" for r in checks["results"]))

    def test_pending_bundle_blocks_final_without_node(self):
        self.ready_review()
        s = lf.current(self.task)
        s["bundle"]["status"] = "pending_business_approval"
        s["bundle_hash"] = digest(s["bundle"])
        lf.commit(self.task, s["revision"], "fixture-pending-policy", s)
        with self.assertRaisesRegex(ValueError, "Final export blocked"):
            self.command("export", "--mode", "final")

    def test_approval_not_inherited_after_content_edit(self):
        self.test_complete_review_then_approval()
        p = fixture_content()
        p["bullets"][0]["text"] = "Three compartments keep pens and small desk items separate"
        self.command("save-content", "--json", self.payload(p))
        self.assertIsNone(lf.current(self.task)["contents"][-1]["approval"])

    def test_literal_formula_text_allowed_as_text_only(self):
        self.ready_content()
        p = fixture_content()
        p["description"][0]["text"] = "=HYPERLINK(\"https://invalid.example\", \"test\")"
        self.command("save-content", "--json", self.payload(p))
        self.assertEqual(lf.current(self.task)["contents"][-1]["content"]["description"][0]["text"][0], "=")

    def test_word_table_and_header_extraction(self):
        from docx import Document
        p = self.source / "table.docx"
        d = Document()
        d.add_paragraph("SKU 000123")
        t = d.add_table(rows=1, cols=1)
        t.cell(0, 0).text = "Desk organizer"
        d.sections[0].header.paragraphs[0].text = "Version A"
        d.save(p)
        self.command("ingest", "--path", str(p))
        texts = [u["text"] for u in lf.current(self.task)["sources"][0]["units"]]
        self.assertIn("Desk organizer", texts)
        self.assertIn("Version A", texts)

    def test_image_requires_observation_for_inspected_gap(self):
        from PIL import Image
        p = self.source / "label.png"
        Image.new("RGB", (12, 12), "white").save(p)
        self.command("ingest", "--path", str(p))
        s = lf.current(self.task)["sources"][0]
        self.assertEqual(s["units"][0]["kind"], "image")
        self.assertTrue(s["gaps"])

    def test_query_invalid_range_rejected(self):
        with self.assertRaisesRegex(ValueError, "offset"):
            self.command("query", "--offset", "-1")

    def test_operator_statement_is_independent_source(self):
        self.ingested()
        self.command("statement", "--json", self.payload({
            "text": "Synthetic batch A", "actor": "TEST", "quote": "Test statement, not an actual business assertion",
            "reason": "Synthetic test"}))
        s = lf.current(self.task)["sources"][-1]
        self.assertEqual(s["units"][0]["kind"], "operator_statement")
        self.assertNotIn("snapshot", s)

    def test_missing_observation_cannot_clear_visual_gap(self):
        self.ingested()
        from PIL import Image
        p = self.source / "label.png"
        Image.new("RGB", (12, 12), "white").save(p)
        self.command("ingest", "--path", str(p))
        facts = self.facts_payload()
        facts["selection"]["S0002"] = {"action": "use", "sku": "000123", "site": "US", "reason": "Test label"}
        facts["coverage"]["S0002:G1"] = {"action": "inspected", "observation_ids": ["made-up"], "reason": "Test"}
        with self.assertRaisesRegex(ValueError, "recorded observations"):
            self.command("freeze", "--json", self.payload(facts))

    def test_actual_observation_clears_only_matching_gap(self):
        self.ingested()
        from PIL import Image
        p = self.source / "label.png"
        Image.new("RGB", (12, 12), "white").save(p)
        self.command("ingest", "--path", str(p))
        self.command("observe", "--json", self.payload({
            "unit_id": "S0002:U1", "method": "visual", "text": "Synthetic white image",
            "actor": "TEST", "reason": "Test observation"}))
        facts = self.facts_payload()
        facts["selection"]["S0002"] = {"action": "use", "sku": "000123", "site": "US", "reason": "Test label"}
        facts["coverage"]["S0002:G1"] = {"action": "inspected", "observation_ids": ["O1"], "reason": "Synthetic matching observation"}
        self.command("freeze", "--json", self.payload(facts))
        self.command("save-content", "--json", self.payload(fixture_content()))
        review = fixture_review(lf.current(self.task))
        with self.assertRaisesRegex(ValueError, "real image review"):
            self.command("review", "--json", self.payload(review))

    def test_pdf_blank_page_still_needs_visual_review(self):
        from pypdf import PdfWriter
        p = self.source / "manual.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(p)
        writer.close()
        self.command("ingest", "--path", str(p))
        s = lf.current(self.task)["sources"][0]
        self.assertEqual(s["status"], "PARTIAL")
        self.assertEqual(s["units"][0]["locator"], "page:1")
        result = self.command("render", "--source", "S0001", "--page", "1")
        self.assertTrue(Path(result["preview"]).is_file())

    def test_modified_original_not_silently_reused(self):
        self.ingested()
        path = self.source / "中文产品资料.xlsx"
        path.write_bytes(b"modified")
        with self.assertRaisesRegex(ValueError, "source changed"):
            self.command("ingest", "--path", str(path))

    def test_plain_python_draft_export_and_leading_zeros(self):
        self.ready_content()
        result = self.command("export", "--mode", "draft")
        from openpyxl import load_workbook
        wb = load_workbook(result["files"][0], read_only=True)
        try:
            self.assertEqual(wb["文案"]["B3"].value, "000123")
            self.assertEqual(wb["文案"]["B3"].data_type, "s")
        finally:
            wb.close()

    def test_specification_name_checked_for_risk_terms(self):
        self.ready_content()
        state = lf.current(self.task)
        state["bundle"]["rules"] = [{
            "id": "R-X", "group": "test", "status": "active", "kind": "text", "terms": ["ForbiddenBrand"],
            "scope": {k: ["*"] for k in ("sku", "site", "language", "category")}, "match": "word_phrase",
            "severity": "blocking", "sources": [], "decision": "Synthetic test rule"}]
        record = state["contents"][-1]
        record["content"]["specs"][0]["name"] = "ForbiddenBrand"
        record["hash"] = digest(record["content"])
        checks = build_checks(state, record)
        self.assertTrue(any(r["id"] == "R-X" and r["status"] == "FAIL" and r["path"] == "specs.0.name"
                            for r in checks["results"]))

    def test_phrase_combination_requires_all_terms(self):
        self.ready_content()
        state = lf.current(self.task)
        state["bundle"]["rules"] = [{
            "id": "R-X", "group": "test", "status": "active", "kind": "combination", "terms": ["Desk", "Nonexistent"],
            "scope": {k: ["*"] for k in ("sku", "site", "language", "category")}, "match": "word_phrase",
            "severity": "blocking", "sources": [], "decision": "Synthetic test rule"}]
        checks = build_checks(state, state["contents"][-1])
        self.assertFalse(any(r["id"] == "R-X" and r["status"] == "FAIL" for r in checks["results"]))

    def test_duplicate_bullets_are_blocked(self):
        self.ready_content()
        content = fixture_content()
        content["bullets"][1] = copy.deepcopy(content["bullets"][0])
        self.command("save-content", "--json", self.payload(content))
        checks = self.command("show", "--section", "checks")
        self.assertTrue(any(r["id"] == "distinct.bullets" and r["status"] == "FAIL" for r in checks["results"]))


class RuleCase(unittest.TestCase):
    def test_word_boundaries_prevent_short_word_false_positives(self):
        self.assertFalse(matches("small glass classic", "all"))
        self.assertFalse(matches("classic", "SS"))
        self.assertTrue(matches("all items", "all"))

    def test_phrase_case_and_spacing(self):
        self.assertTrue(matches("NO   MORE", "No More"))
        self.assertFalse(matches("NoMore", "No More"))

    def test_scope_unknown_is_not_universal(self):
        rule = {"scope": {"sku": ["A"], "site": None, "language": ["*"], "category": ["*"]}}
        self.assertEqual(scope_result(rule, {"sku": "A", "site": "US"}), "unknown")
        self.assertEqual(scope_result(rule, {"sku": "B", "site": "US"}), "no")

    def test_eu_not_silently_de(self):
        rule = {"scope": {"sku": ["*"], "site": ["EU"], "language": ["*"], "category": ["*"]}}
        self.assertEqual(scope_result(rule, {"site": "DE"}), "no")

    def test_pending_policy_cannot_be_published(self):
        b = approved_bundle()
        b["policy"]["pending_decisions"] = ["unknown"]
        with self.assertRaisesRegex(ValueError, "remain unresolved"):
            validate_bundle(b)

    def test_unsupported_matcher_rejected(self):
        b = approved_bundle()
        b["rules"] = [{"id": "R-1", "kind": "text", "match": "regex", "status": "active",
                       "severity": "blocking", "terms": ["word"],
                       "scope": {k: ["*"] for k in ("sku", "site", "language", "category")}}]
        with self.assertRaisesRegex(ValueError, "Unsupported matcher"):
            validate_bundle(b)

    def test_xml_entity_rejected(self):
        with self.assertRaisesRegex(ValueError, "DTD"):
            xml(b'<!DOCTYPE x [<!ENTITY a "evil">]><x>&a;</x>')

    def test_active_eu_scope_requires_explicit_mapping(self):
        b = approved_bundle()
        b["rules"] = [{"id": "R-1", "kind": "text", "match": "word_phrase", "status": "active",
                       "severity": "blocking", "terms": ["word"], "decision": "Test",
                       "scope": {"sku": ["*"], "site": ["EU"], "language": ["*"], "category": ["*"]}}]
        with self.assertRaisesRegex(ValueError, "resolve EU"):
            validate_bundle(b)

    def test_real_bundle_provenance_and_quarantine(self):
        b = lf.load(lf.SKILL / "assets/bundle.json")
        self.assertEqual(len(b["rules"]), 340)
        self.assertEqual(len(b["quarantine"]), 1)
        self.assertTrue(all(r["sources"][0]["sha256"] for r in b["rules"]))
        self.assertEqual(b["status"], "pending_business_approval")


if __name__ == "__main__":
    unittest.main()
