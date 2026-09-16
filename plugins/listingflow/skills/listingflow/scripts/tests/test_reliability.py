"""V3 contracts, failure-shaped regressions, and operator export checks."""

import unittest
from pathlib import Path

import test_listingflow as base
import quality_fixtures as fixture
import listingflow as lf
import lf_reliability as reliable
from lf_excel import tables


def reconcile(state, payload):
    units = {u["id"]: u for s in state["sources"] for u in s["units"]}
    payload["source_review"] = {s["id"]: {"role": "primary", "reviewed": True,
                                       "reason": "合成测试文字来源，已读取适用字段。"} for s in state["sources"]}
    payload["reconciliation"] = [
        {"field": f["field"], "status": "adopted", "fact_ids": [f["id"]], "critical": f["id"] == "F7",
         "candidates": [{"source_id": sid, "quote": units[sid]["text"], "value": f["value"]} for sid in f["sources"]],
         "reason": "合成资料明确写明此字段；没有其他候选值。", "action": "采用并保留条件。", "owner": "文案助手"}
        for f in payload["facts"]]
    return payload


def claim_review(state):
    review = fixture.review(state)
    blocks = dict(lf.content_blocks(state["contents"][-1]["content"]))
    for path, block in blocks.items():
        review["blocks"][path]["reason"] = "合成验收样例：已对照所列字段检查原句，不是实际商品核对结论。"
        review["blocks"][path]["claims"] = [
            {"text": block["text"], "fact_ids": block["fact_ids"], "status": "PASS",
             "reason": "合成测试：所列事实支持该句；不表示真实产品通过审核。"}]
    review["conditions"] = {
        fid: {"path": path, "quote": block["text"], "reason": "合成测试核对清洁条件。"}
        for fid in state["brief"]["critical_fact_ids"]
        for path, block in blocks.items() if fid in block["fact_ids"]}
    for item in review["quality"].values():
        item["reason"] = "合成验收样例：检查当前例句的表达，不代表真实运营质量评定。"
    from lf_checklist import ITEMS, content_texts
    review["operator_checklist"] = {
        key: {"status": "NOT_CHECKED" if key in {"ip.visual", "local.search_data"} else "PASS",
              "paths": list(content_texts(state["contents"][-1]["content"])),
              "examples": [{"path": "title", "quote": blocks["title"]["text"]}],
              "reason": "合成四类自查样例，仅用于验证程序，不是实际商品审核。",
              "action": "合成测试无待办。"} for key in ITEMS}
    return review


class ReliabilityCase(unittest.TestCase):
    def setUp(self):
        self.flow = base.FlowCase()
        self.flow.use_v3 = True
        self.flow.setUp()

    def tearDown(self):
        self.flow.tearDown()

    def state(self):
        return lf.current(self.flow.task)

    def send(self, command, value):
        return self.flow.command(command, "--json", self.flow.payload(value))

    def facts(self):
        self.flow.ingested()
        return reconcile(self.state(), self.flow.facts_payload())

    def ready(self, payload=None):
        self.send("freeze", payload or self.facts())
        self.send("research", fixture.research(self.state()))
        self.send("brief", fixture.brief(self.state()))
        self.send("save-content", base.fixture_content())

    def test_new_tasks_use_v3(self):
        self.assertEqual(self.state()["workflow_version"], 3)

    def test_fact_ledger_required(self):
        self.flow.ingested()
        with self.assertRaisesRegex(ValueError, "source_review"):
            self.send("freeze", self.flow.facts_payload())

    def test_missing_field_rejected(self):
        payload = self.facts()
        payload["reconciliation"].pop()
        with self.assertRaisesRegex(ValueError, "exactly once"):
            self.send("freeze", payload)

    def test_weight_conflict_keeps_other_fields(self):
        payload = self.facts()
        payload["reconciliation"].append({
            "field": "净重", "status": "omitted", "fact_ids": [], "critical": False, "candidates": [],
            "reason": "测试场景：重量有差异；此处没有真实摘录，不冒充核验结果。",
            "action": "本版不写重量，其他字段继续采用。", "owner": "产品负责人"})
        self.ready(payload)
        checks = lf.build_checks(self.state(), self.state()["contents"][-1])
        row = next(r for r in checks["results"] if r["id"] == "reliability.field.净重")
        self.assertEqual(row["severity"], "info")
        self.assertEqual(len(self.state()["facts"]["facts"]), 7)

    def test_conflict_cannot_be_confirmed(self):
        payload = self.facts()
        payload["reconciliation"][3]["status"] = "pending"
        with self.assertRaisesRegex(ValueError, "cannot enter"):
            self.send("freeze", payload)

    def test_critical_omission_blocks(self):
        payload = self.facts()
        payload["reconciliation"].append({
            "field": "安全限制", "status": "pending", "fact_ids": [], "critical": True, "candidates": [],
            "reason": "缺少关键条件。", "action": "先确认。", "owner": "产品负责人"})
        with self.assertRaisesRegex(ValueError, "critical field"):
            self.send("freeze", payload)

    def test_manual_cannot_be_skipped(self):
        payload = self.facts()
        payload["source_review"]["S0001"].update(role="manual", reviewed=False)
        with self.assertRaisesRegex(ValueError, "manuals must be reviewed"):
            self.send("freeze", payload)

    def test_manual_needs_visual_observations(self):
        payload = self.facts()
        state = self.state()
        state["sources"][0]["units"][0]["kind"] = "image"
        payload["source_review"]["S0001"]["role"] = "manual"
        with self.assertRaisesRegex(ValueError, "visual observations"):
            reliable.validate_facts(payload, state)

    def test_fabricated_ledger_quote_rejected(self):
        payload = self.facts()
        payload["reconciliation"][0]["candidates"][0]["quote"] = "not present"
        with self.assertRaisesRegex(ValueError, "does not exist"):
            self.send("freeze", payload)

    def test_custom_field_requires_operator_label(self):
        payload = self.facts()
        payload["facts"][0]["field"] = "custom_unknown_field"
        payload["reconciliation"][0]["field"] = "custom_unknown_field"
        with self.assertRaisesRegex(ValueError, "Chinese operator label"):
            reliable.validate_facts(payload, self.state())

    def test_dimension_candidates_preserved_without_arbitrary_selection(self):
        payload = self.facts()
        state = self.state()
        source = state["sources"][0]
        source["units"].extend([
            {"id": "S0001:U90", "text": "21 x 9 x 2.5 cm", "kind": "cell", "locator": "Sheet1!C6"},
            {"id": "S0001:U91", "text": "20 x 9 cm", "kind": "cell", "locator": "Manual!A2"}])
        row = payload["reconciliation"][3]
        row.update(status="pending", fact_ids=[], candidates=[
            {"source_id": "S0001:U90", "value": "21 x 9 x 2.5 cm", "quote": "21 x 9 x 2.5 cm"},
            {"source_id": "S0001:U91", "value": "20 x 9 cm", "quote": "20 x 9 cm"}],
            reason="两个测试来源的尺寸不同，未确认测量口径。", action="本版不写尺寸，确认版本后再采用。")
        payload["facts"] = [f for f in payload["facts"] if f["id"] != "F4"]
        reliable.validate_facts(payload, state)
        self.assertEqual(len(row["candidates"]), 2)

    def test_unrelated_image_observation_cannot_cover_fact(self):
        self.ready()
        state = self.state()
        entry = state["facts"]["reconciliation"][0]["candidates"][0]
        next(u for u in state["sources"][0]["units"] if u["id"] == entry["source_id"])["kind"] = "image"
        entry["observation_id"] = "O1"
        state["observations"] = [
            {"id": "O1", "unit_id": entry["source_id"], "method": "visual", "text": entry["quote"]},
            {"id": "O2", "unit_id": "S0001:U1", "method": "visual", "text": "irrelevant image"}]
        review = claim_review(state)
        review["fact_evidence"]["F1"][0]["observation_id"] = "O1"
        review["categories"]["images"].update(status="PASS", observation_ids=["O2"])
        with self.assertRaisesRegex(ValueError, "cover visual"):
            reliable.validate_review(review, state, state["contents"][-1])

    def test_brief_preserves_critical_condition(self):
        self.send("freeze", self.facts())
        brief = fixture.brief(self.state())
        brief["critical_fact_ids"] = []
        with self.assertRaisesRegex(ValueError, "critical facts"):
            self.send("brief", brief)

    def test_v2_review_insufficient(self):
        self.ready()
        with self.assertRaisesRegex(ValueError, "claim-level"):
            self.send("review", fixture.review(self.state()))

    def test_uncovered_claim_rejected(self):
        self.ready()
        review = claim_review(self.state())
        review["blocks"]["title"]["claims"][0]["text"] = "Desk"
        with self.assertRaisesRegex(ValueError, "uncovered"):
            self.send("review", review)

    def test_unsupported_claim_cannot_pass(self):
        self.ready()
        review = claim_review(self.state())
        review["blocks"]["title"]["claims"][0]["status"] = "NEEDS_REVIEW"
        with self.assertRaisesRegex(ValueError, "cannot pass"):
            self.send("review", review)

    def test_condition_must_match_content(self):
        self.ready()
        review = claim_review(self.state())
        review["conditions"]["F7"]["quote"] = "Machine washable"
        with self.assertRaisesRegex(ValueError, "Condition must"):
            self.send("review", review)

    def test_visual_pass_requires_observation(self):
        self.ready()
        review = claim_review(self.state())
        review["categories"]["images"]["status"] = "PASS"
        with self.assertRaisesRegex(ValueError, "visual observations"):
            self.send("review", review)

    def test_complete_review_allows_fixture_approval(self):
        self.ready()
        self.send("review", claim_review(self.state()))
        self.send("approve", {"content_hash": self.state()["contents"][-1]["hash"],
                              "actor": "SYNTHETIC TEST", "quote": "Synthetic fixture approval only"})
        self.assertTrue(self.state()["contents"][-1]["approval"])

    def guard_results(self, claim, quote):
        self.ready()
        state = self.state()
        record = state["contents"][-1]
        record["content"]["description"][0] = {"text": claim, "fact_ids": ["F1"]}
        review = claim_review(state)
        review["fact_evidence"]["F1"][0]["quote"] = quote
        record["review"] = review
        return reliable.checks(state, record)

    def test_voltage_cannot_prove_battery(self):
        results = self.guard_results("USB corded power with no built-in battery.", "5V/1.6A; 1.5m cable")
        self.assertTrue(any(r["id"].startswith("reliability.battery") for r in results))

    def test_direct_battery_evidence_not_flagged(self):
        results = self.guard_results("No built-in battery.", "USB插电（不带电池）")
        self.assertFalse(any(r["id"].startswith("reliability.battery") for r in results))

    def test_fit_cannot_prove_strap(self):
        results = self.guard_results("An adjustable strap positions the mask.", "3D contour fits around the eyes")
        self.assertTrue(any(r["id"].startswith("reliability.strap") for r in results))

    def test_recorded_sliding_buckle_observation_supports_strap(self):
        results = self.guard_results("An adjustable strap lets you change the fit.", "绑带可通过滑扣调节长度")
        self.assertFalse(any(r["id"].startswith("reliability.strap") for r in results))

    def test_unrelated_rationale_cannot_prove_controls(self):
        results = self.guard_results("Heat and vibration can be adjusted separately.", "未扩大为医疗模式")
        self.assertTrue(any(r["id"].startswith("reliability.independent") for r in results))

    def test_irrelevant_reason_flagged_even_with_control_evidence(self):
        self.ready()
        state = self.state()
        record = state["contents"][-1]
        record["content"]["description"][0] = {"text": "Heat and vibration can be adjusted separately.", "fact_ids": ["F1"]}
        record["review"] = claim_review(state)
        record["review"]["fact_evidence"]["F1"][0]["quote"] = "Heat and vibration have independent controls."
        record["review"]["blocks"]["description.0"]["claims"][0]["reason"] = "未扩大为医疗模式。"
        self.assertTrue(any(r["id"].startswith("reliability.rationale") for r in reliable.checks(state, record)))

    def test_timer_reset_not_omitted(self):
        results = self.guard_results("A 15-minute timer turns the mask off.", "15-minute timer resets with the last button press.")
        self.assertTrue(any(r["id"].startswith("reliability.timer") for r in results))

    def test_restarts_is_a_valid_timer_reset_paraphrase(self):
        results = self.guard_results("Auto-off after 15 minutes; each button press restarts the timer.",
                                     "The timer resets upon the last button press.")
        self.assertFalse(any(r["id"].startswith("reliability.timer") for r in results))

    def test_digital_manual_not_physical(self):
        results = self.guard_results("Package includes a multilingual manual.", "英文实物说明书；八国语言电子档说明书")
        self.assertTrue(any(r["id"].startswith("reliability.package") for r in results))

    def test_temperature_tolerance_units(self):
        results = self.guard_results("95°F to 131°F; tolerance ±5°C.", "35°C to 55°C; tolerance ±5°C")
        self.assertTrue(any(r["id"].startswith("reliability.units") for r in results))

    def exported(self):
        self.ready()
        self.send("review", claim_review(self.state()))
        result = self.flow.command("export", "--mode", "draft")
        return result, lf.load(Path(result["files"][0]).parent / "export.json")

    def test_operator_export_and_full_audit(self):
        result, request = self.exported()
        from openpyxl import load_workbook
        self.assertEqual(request["checks"], lf.build_checks(self.state(), self.state()["contents"][-1]))
        wb = load_workbook(result["files"][1])
        try:
            self.assertEqual(wb.sheetnames, ["审核结论", "违禁词自查表", "侵权词自查表", "本地化自查表", "文案合规自查表",
                                             "处理清单", "已处理与可选", "内容检查", "来源依据", "规则与范围", "版本"])
            self.assertEqual(wb["违禁词自查表"].freeze_panes, "C2")
            self.assertEqual(wb["处理清单"]["F1"].value, "处理意见")
            self.assertEqual(wb["版本"].sheet_state, "hidden")
            self.assertTrue(wb["内容检查"].auto_filter.ref)
            self.assertTrue(wb["处理清单"].data_validations.dataValidation)
            self.assertEqual(wb.active.title, "审核结论")
            self.assertTrue(wb["审核结论"].protection.sheet)
            self.assertTrue(wb["审核结论"]["B4"].protection.locked)
            self.assertFalse(wb["处理清单"]["F3"].protection.locked)
            self.assertTrue(wb["处理清单"]["D3"].protection.locked)
            self.assertEqual(wb["审核结论"]["F2"].hyperlink.location, "'内容检查'!A1")
            self.assertEqual(wb["处理清单"].data_validations.dataValidation[0].sqref, f"I3:I{wb['处理清单'].max_row}")
        finally:
            wb.close()
        wb = load_workbook(result["files"][0])
        try:
            self.assertEqual(wb["文案"].max_column, 2)
            self.assertEqual(wb["文案"]["B3"].value, "000123")
        finally:
            wb.close()

    def test_pending_business_rules_block_final(self):
        self.exported()
        state = self.state()
        state["bundle"] = lf.load(lf.SKILL / "assets/bundle.json")
        state["bundle_hash"] = lf.digest(state["bundle"])
        lf.commit(self.flow.task, state["revision"], "pending-v3-fixture", state)
        with self.assertRaisesRegex(ValueError, "Final export blocked"):
            self.flow.command("export", "--mode", "final")
        result = self.flow.command("export", "--mode", "draft")
        request = lf.load(Path(result["files"][0]).parent / "export.json")
        rows = next(s[2] for s in tables(request)["自查"] if s[0] == "处理清单")
        self.assertTrue(any(r[0].startswith("公司规则确认") for r in rows))
        self.assertFalse(any("R-0318" in str(r) for r in rows))

    def test_template_contains_no_fabricated_passes(self):
        self.ready()
        template = lf.review_template(self.state())
        self.assertTrue(all(not b["claims"] for b in template["blocks"].values()))
        self.assertEqual(template["conditions"]["F7"]["quote"], "")

    def test_problem_overrides_pass_in_content_sheet(self):
        _, request = self.exported()
        request["checks"]["results"].append({"id": "reliability.test", "path": "title", "status": "NEEDS_REVIEW",
                                             "severity": "blocking", "reason": "实际依据待核对"})
        sheet = next(s for s in tables(request)["自查"] if s[0] == "内容检查")
        self.assertEqual(sheet[2][0][2], "待核实")

    def test_old_export_request_retains_original_sheet_layout(self):
        _, request = self.exported()
        request.pop("report_layout")
        request.pop("rule_catalog")
        self.assertEqual([s[0] for s in tables(request)["自查"]],
                         ["处理清单", "内容检查", "来源依据", "规则与范围", "版本"])

    def test_old_v3_checks_not_silently_upgraded(self):
        self.ready()
        state = self.state()
        state.pop("checklist_version")
        review = claim_review(state)
        review.pop("operator_checklist")
        state["contents"][-1]["review"] = review
        checks = lf.build_checks(state, state["contents"][-1])
        self.assertFalse(any(r["id"].startswith("checklist.") for r in checks["results"]))
        self.assertTrue(checks["ready_for_human_approval"])

    def test_specific_failure_visible_in_content_check(self):
        _, request = self.exported()
        request["checks"]["results"].append({
            "id": "checklist.compliance.medical", "status": "FAIL", "severity": "blocking",
            "path": "checklist.compliance.medical", "reason": "合成医疗宣称失败。"})
        sheet = next(s for s in tables(request)["自查"] if s[0] == "内容检查")
        self.assertEqual(sheet[2][0][2], "需修改")

    def test_manual_rule_review_not_counted_as_exact_scan(self):
        _, request = self.exported()
        request["checks"]["results"].append({
            "id": "R-MANUAL", "status": "PASS", "severity": "blocking", "path": "rules",
            "reason": "合成专项核验。", "method": "explicit_rule_review"})
        rows = next(s[2] for s in tables(request)["自查"] if s[0] == "规则与范围")
        self.assertTrue(next(r[2] for r in rows if r[0] == "精确风险词扫描").startswith("0条"))
        self.assertTrue(any(r[0] == "专项规则核验" for r in rows))

    def test_visual_placeholder_not_evidence(self):
        self.ready()
        state = self.state()
        unit = state["sources"][0]["units"][1]
        unit["kind"] = "image"
        with self.assertRaisesRegex(ValueError, "recorded observation"):
            reliable.excerpt({"source_id": unit["id"], "quote": unit["text"]}, state)
