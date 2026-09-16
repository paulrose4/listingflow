"""V2 workflow, evidence, operator experience, and compatibility regression."""

import copy
import unittest

import test_listingflow as base
import quality_fixtures as fixture
import listingflow as lf
from lf_experience import dashboard


class EditorialCase(unittest.TestCase):
    def setUp(self):
        self.flow = base.FlowCase()
        self.flow.use_v2 = True
        self.flow.setUp()

    def tearDown(self):
        self.flow.tearDown()

    def state(self):
        return lf.current(self.flow.task)

    def send(self, name, payload):
        return self.flow.command(name, "--json", self.flow.payload(payload))

    def ready(self):
        fixture.ready_content(self.flow)

    def reviewed(self):
        self.ready()
        self.send("review", fixture.review(self.state()))

    def approved(self):
        self.reviewed()
        self.send("approve", {"content_hash": self.state()["contents"][-1]["hash"], "actor": "TEST",
                              "quote": "Synthetic test confirmation, not real business approval"})

    def test_new_tasks_snapshot_editorial_resource(self):
        self.assertEqual(self.state()["workflow_version"], 2)
        self.assertEqual(self.state()["editorial_hash"], lf.digest(self.state()["editorial"]))

    def test_brief_required_for_v2(self):
        self.flow.ingested()
        self.send("freeze", self.flow.facts_payload())
        with self.assertRaisesRegex(ValueError, "writing brief"):
            self.send("save-content", base.fixture_content())

    def test_duplicate_angles_rejected(self):
        self.flow.ingested()
        self.send("freeze", self.flow.facts_payload())
        brief = fixture.brief(self.state())
        brief["angles"][1]["label"] = brief["angles"][0]["label"]
        with self.assertRaisesRegex(ValueError, "Duplicate angles"):
            self.send("brief", brief)

    def test_cross_task_brief_rejected(self):
        self.ready()
        brief = fixture.brief(self.state())
        brief["facts_hash"] = "other-product"
        with self.assertRaisesRegex(ValueError, "frozen task"):
            self.send("brief", brief)

    def test_unsupported_brief_fact_rejected(self):
        self.ready()
        brief = fixture.brief(self.state())
        brief["angles"][0]["fact_ids"] = ["imaginary"]
        with self.assertRaisesRegex(ValueError, "fact reference"):
            self.send("brief", brief)

    def test_old_style_review_not_enough_for_v2(self):
        self.ready()
        with self.assertRaisesRegex(ValueError, "brief hash"):
            self.send("review", base.fixture_review(self.state()))

    def test_invented_evidence_quote_rejected(self):
        self.ready()
        review = fixture.review(self.state())
        review["fact_evidence"]["F1"][0]["quote"] = "This assertion does not exist in the original"
        with self.assertRaisesRegex(ValueError, "does not exist"):
            self.send("review", review)

    def test_fact_cannot_quote_unrelated_source(self):
        self.ready()
        review = fixture.review(self.state())
        review["fact_evidence"]["F1"] = copy.deepcopy(review["fact_evidence"]["F2"])
        with self.assertRaisesRegex(ValueError, "not attached"):
            self.send("review", review)

    def test_invented_quality_example_rejected(self):
        self.ready()
        review = fixture.review(self.state())
        review["quality"]["naturalness"]["examples"][0]["quote"] = "This text was not generated"
        with self.assertRaisesRegex(ValueError, "actual current content"):
            self.send("review", review)

    def test_quality_pass_without_example_rejected(self):
        self.ready()
        review = fixture.review(self.state())
        review["quality"]["conciseness"]["examples"] = []
        with self.assertRaisesRegex(ValueError, "concrete text examples"):
            self.send("review", review)

    def test_complete_evidence_review_allows_confirmation(self):
        self.approved()
        self.assertTrue(self.state()["contents"][-1]["approval"])

    def test_brief_change_invalidates_approval_but_preserves_history(self):
        self.approved()
        previous = self.state()["contents"][-1]["brief_hash"]
        brief = fixture.brief(self.state())
        brief["tone"] = "technical_precise"
        self.send("brief", brief)
        self.assertIsNone(self.state()["contents"][-1]["approval"])
        self.assertEqual(self.state()["contents"][-1]["brief_hash"], previous)
        self.assertNotEqual(self.state()["brief_hash"], previous)
        with self.assertRaisesRegex(ValueError, "Final export blocked"):
            self.flow.command("export", "--mode", "final")

    def test_partial_patch_preserves_other_blocks(self):
        self.reviewed()
        before = self.state()["contents"][-1]["content"]
        change = {"bullets.2": {"text": "Open tops provide access to the stored items", "fact_ids": ["F5"]}}
        result = self.send("revise", change)
        after = self.state()["contents"][-1]
        self.assertEqual(after["content"]["title"], before["title"])
        self.assertEqual(after["content"]["qa"], before["qa"])
        self.assertEqual(result["changed_paths"], ["bullets.2"])
        self.assertIsNone(after["review"])

    def test_unchanged_patch_keeps_approval_and_revision(self):
        self.approved()
        before = self.state()
        result = self.send("revise", {"title": before["contents"][-1]["content"]["title"]})
        self.assertTrue(result["unchanged"])
        self.assertEqual(self.state()["revision"], before["revision"])
        self.assertEqual(self.state()["contents"][-1]["approval"], before["contents"][-1]["approval"])

    def test_unchanged_brief_keeps_approval_and_revision(self):
        self.approved()
        before = self.state()
        result = self.send("brief", before["brief"])
        self.assertTrue(result["unchanged"])
        self.assertEqual(self.state()["revision"], before["revision"])
        self.assertEqual(self.state()["contents"][-1]["approval"], before["contents"][-1]["approval"])

    def test_patch_cannot_mutate_unlisted_properties(self):
        self.ready()
        with self.assertRaisesRegex(ValueError, "Unknown partial rewrite path"):
            self.send("revise", {"../../bundle": {"status": "approved"}})

    def test_operator_lock_requires_explicit_unlock(self):
        self.ready()
        self.send("lock", {"action": "lock", "paths": ["title"], "actor": "TEST", "quote": "Fixture: keep title"})
        with self.assertRaisesRegex(ValueError, "locked block"):
            self.send("revise", {"title": {"text": "Desk Organizer with Open Compartments", "fact_ids": ["F1", "F5"]}})
        self.send("lock", {"action": "unlock", "paths": ["title"], "actor": "TEST", "quote": "Fixture: allow title edit"})
        self.send("revise", {"title": {"text": "Desk Organizer with Open Compartments", "fact_ids": ["F1", "F5"]}})

    def test_critical_fact_omission_flagged(self):
        self.ready()
        self.send("revise", {"bullets.4": {"text": "A desk organizer for pens and small items", "fact_ids": ["F2"]},
                             "qa.4.question": {"text": "What can it hold?", "fact_ids": ["F2"]},
                             "qa.4.answer": {"text": "Pens and small items.", "fact_ids": ["F2"]}})
        checks = self.flow.command("show", "--section", "checks")
        self.assertTrue(any(r["id"] == "editorial.critical_facts" and r["status"] == "FAIL" for r in checks["results"]))

    def test_near_duplicate_detected_as_warning_not_verdict(self):
        self.ready()
        self.send("revise", {
            "bullets.0": {"text": "Three open compartments neatly separate small desk items", "fact_ids": ["F2", "F3", "F5"]},
            "bullets.1": {"text": "Three open compartments separate small desk items neatly", "fact_ids": ["F2", "F3", "F4", "F5"]}})
        checks = self.flow.command("show", "--section", "checks")
        self.assertTrue(any(r["id"].startswith("editorial.overlap") and r["severity"] == "warning" for r in checks["results"]))

    def test_number_conversion_is_not_automatically_blocked(self):
        self.ready()
        self.send("revise", {"specs.1": {"name": "Dimensions", "text": "Approx. 7.1 x 3.9 x 3.5 in", "fact_ids": ["F4"]}})
        checks = self.flow.command("show", "--section", "checks")
        warning = next(r for r in checks["results"] if r["id"] == "editorial.numbers.specs.1")
        self.assertEqual(warning["severity"], "warning")
        self.assertEqual(self.state()["contents"][-1]["content"]["specs"][1]["text"], "Approx. 7.1 x 3.9 x 3.5 in")

    def test_english_filler_warning(self):
        self.ready()
        self.send("revise", {"description.0": {"text": "The ultimate solution for desk organization", "fact_ids": ["F1"]}})
        checks = self.flow.command("show", "--section", "checks")
        self.assertTrue(any(r["id"].startswith("editorial.filler") for r in checks["results"]))

    def test_german_filler_warning_is_language_specific(self):
        self.ready()
        state = self.state()
        state["language"] = "de"
        record = state["contents"][-1]
        record["content"]["description"][0]["text"] = "Die ultimative Lösung für Ihren Schreibtisch."
        checks = lf.build_checks(state, record)
        self.assertTrue(any(r["id"].startswith("editorial.filler") for r in checks["results"]))

    def test_dashboard_has_no_passing_rule_dump(self):
        self.ready()
        view = self.flow.command("dashboard")
        self.assertEqual(view["next_owner"], "文案助手")
        self.assertNotIn("results", view)
        self.assertTrue(all(len(g["examples"]) <= 3 for g in view["issues"]))

    def test_rule_failure_is_assigned_to_copy_assistant(self):
        self.reviewed()
        checks = {"results": [{"id": "R-test", "status": "FAIL", "severity": "blocking",
                              "path": "title", "reason": "Synthetic prohibited term match"}],
                  "blocking_count": 1}
        view = dashboard(self.state(), checks)
        self.assertEqual(view["next_owner"], "文案助手")
        self.assertEqual(view["issues"][0]["group"], "fix_copy")

    def test_business_confirmation_is_grouped_not_silently_passed(self):
        self.reviewed()
        state = self.state()
        state["bundle"] = lf.load(lf.SKILL / "assets/bundle.json")
        state["bundle_hash"] = lf.digest(state["bundle"])
        lf.commit(self.flow.task, state["revision"], "pending-business-test-fixture", state)
        view = self.flow.command("dashboard")
        group = next(g for g in view["issues"] if g["group"] == "business_rules")
        self.assertGreater(group["count"], 1)
        self.assertEqual(group["owner"], "规则负责人")
        with self.assertRaisesRegex(ValueError, "Final export blocked"):
            self.flow.command("export", "--mode", "final")

    def test_prepare_does_not_invent_review_passes(self):
        self.ready()
        revision = self.state()["revision"]
        result = self.flow.command("prepare", "--kind", "review")
        template = lf.load(result["template"])
        self.assertTrue(all(i["status"] == "NOT_CHECKED" for i in template["quality"].values()))
        self.assertEqual(self.state()["revision"], revision)
        self.assertTrue(all(not v for v in template["fact_evidence"].values()))

    def test_evidence_is_source_excerpt_not_a_new_fact(self):
        self.ready()
        result = self.flow.command("evidence", "--fact", "F3")
        self.assertEqual(result["evidence"][0]["text"], "3")
        self.assertIn("!A4", result["evidence"][0]["locator"])

    def test_context_excludes_raw_source_prompts(self):
        self.ready()
        context = self.flow.command("context")
        self.assertNotIn("source_prompts", context)
        self.assertNotIn("sources", context)
        self.assertEqual(context["brief_hash"], self.state()["brief_hash"])

    def test_task_lookup_preserves_exact_sku(self):
        result = lf.run(lf.parser().parse_args(["tasks", "--root", str(self.flow.root / "tasks"), "--sku", "000123"]))
        self.assertEqual(len(result["tasks"]), 1)
        result = lf.run(lf.parser().parse_args(["tasks", "--root", str(self.flow.root / "tasks"), "--sku", "123"]))
        self.assertEqual(result["tasks"], [])

    def test_editorial_snapshot_tamper_blocks_resume(self):
        state = self.state()
        state["editorial"]["near_duplicate_threshold"] = 1.0
        lf.commit(self.flow.task, state["revision"], "tamper-fixture", state)
        with self.assertRaisesRegex(ValueError, "Editorial resource hash"):
            self.flow.command("dashboard")

    def test_summary_export_contains_action_sheet(self):
        self.reviewed()
        result = self.flow.command("export", "--mode", "draft")
        from openpyxl import load_workbook
        wb = load_workbook(result["files"][1], read_only=True)
        try:
            self.assertEqual(wb.sheetnames[0], "处理清单")
            self.assertIn("自查", wb.sheetnames)
            self.assertIn("依据", wb.sheetnames)
            evidence = list(wb["依据"].values)
            self.assertTrue(any(row[0] == "审核摘录" for row in evidence))
            self.assertEqual(sum(row[0] == "文案审核" for row in evidence), 6)
        finally:
            wb.close()


class LegacyCompatibilityCase(unittest.TestCase):
    def test_legacy_approval_survives_check_without_new_editorial_gate(self):
        flow = base.FlowCase()
        flow.setUp()
        try:
            flow.test_complete_review_then_approval()
            before = lf.current(flow.task)["contents"][-1]["approval"]
            flow.command("check")
            after = lf.current(flow.task)["contents"][-1]
            self.assertEqual(before, after["approval"])
            self.assertEqual(after["checks"]["engine"], "listingflow-checks-1")
            self.assertFalse(any(r["id"].startswith("editorial.") for r in after["checks"]["results"]))
        finally:
            flow.tearDown()
