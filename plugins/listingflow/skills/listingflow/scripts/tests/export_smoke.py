"""Isolated synthetic export acceptance; never approves the installed company bundle."""

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from test_listingflow import FlowCase, fixture_content
import listingflow as lf
import quality_fixtures
from test_reliability import reconcile, claim_review


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--output", required=True)
    p.add_argument("--enhanced", action="store_true")
    p.add_argument("--operator", action="store_true")
    args = p.parse_args()
    output = Path(args.output).resolve()
    output.mkdir(parents=True, exist_ok=False)
    case = FlowCase()
    case.use_v2 = args.enhanced
    case.use_v3 = args.operator
    case.setUp()
    try:
        if args.operator:
            case.ingested()
            payload = reconcile(lf.current(case.task), case.facts_payload())
            payload["reconciliation"].append({
                "field": "净重", "status": "omitted", "fact_ids": [], "critical": False, "candidates": [],
                "reason": "合成演示：重量存在未解决差异，本版不写；不是真实商品核对结论。",
                "action": "本版可不处理；需要补写重量时由产品负责人确认。", "owner": "产品负责人"})
            case.command("freeze", "--json", case.payload(payload))
            from test_research import collected
            case.command("research", "--json", case.payload(collected(lf.current(case.task), pla=False)))
            brief = quality_fixtures.brief(lf.current(case.task))
            brief["primary_keyword"].update(source="ebay_research", reference="合成测试：下拉与竞品已取得，PLA访问受限。")
            case.command("brief", "--json", case.payload(brief))
            content = fixture_content()
            content["keywords"] = brief["primary_keyword"]
            case.command("save-content", "--json", case.payload(content))
            case.command("review", "--json", case.payload(claim_review(lf.current(case.task))))
        elif args.enhanced:
            quality_fixtures.ready_content(case)
            case.command("review", "--json", case.payload(quality_fixtures.review(lf.current(case.task))))
        else:
            case.ready_review()
        state = lf.current(case.task)
        case.command("approve", "--json", case.payload({
            "content_hash": state["contents"][-1]["hash"], "actor": "SYNTHETIC ACCEPTANCE FIXTURE",
            "quote": "This is a fixture approval for software testing only, not a real business approval."}))
        result = case.command("export", "--mode", "final", "--preview")
        expected_previews = 14 if args.operator else 6 if args.enhanced else 5
        preview_count = len(list(Path(result["files"][0]).parent.glob("*.png")))
        assert preview_count == expected_previews, result["preview"]
        from openpyxl import load_workbook
        wb = load_workbook(result["files"][0], read_only=True, data_only=False)
        try:
            rows = list(wb["文案"].values)
            assert next(r[1] for r in rows if r[0] == "SKU") == "000123"
            assert next(r[1] for r in rows if r[0] == "标题") == fixture_content()["title"]["text"]
        finally:
            wb.close()
        import shutil
        shutil.copytree(Path(result["files"][0]).parent, output / "approved-fixture")
        content = fixture_content()
        malicious = '=HYPERLINK("https://invalid.example","not executed")'
        content["description"][0]["text"] = malicious
        case.command("save-content", "--json", case.payload(content))
        draft = case.command("export", "--mode", "draft", "--preview")
        draft_preview_count = len(list(Path(draft["files"][0]).parent.glob("*.png")))
        assert draft_preview_count == expected_previews, draft["preview"]
        wb = load_workbook(draft["files"][0], read_only=True, data_only=False)
        try:
            cell = next(c for row in wb["文案"] for c in row if isinstance(c.value, str) and "HYPERLINK" in c.value)
            assert cell.data_type != "f"
            assert cell.value.lstrip("'") == malicious
            assert any("草稿" in str(c.value) for row in wb["文案"] for c in row)
        finally:
            wb.close()
        shutil.copytree(Path(draft["files"][0]).parent, output / "draft-fixture")
        lf.write_new(output / "acceptance.json", {
            "scope": "Synthetic fixture only; installed company bundle remains unapproved",
            "final_pair": "verified", "draft_pair": "verified", "leading_zero_sku": "preserved",
            "literal_formula": "not executed", "stale_review": "cleared on new content",
            "preview_sheets": preview_count, "draft_preview_sheets": draft_preview_count,
            "preview_runtime": {"final": result["preview"], "draft": draft["preview"]},
            "workflow": 3 if args.operator else 2 if args.enhanced else 1,
            "created": lf.now()})
        print(json.dumps({"ok": True, "output": str(output)}, ensure_ascii=False))
    except Exception:
        import shutil
        if (case.task / "exports").exists():
            shutil.copytree(case.task / "exports", output / "failed-export")
        raise
    finally:
        case.tearDown()


if __name__ == "__main__":
    main()
